"""Build a roadshow deck (.pptx) from a JSON slide spec in the brand's fonts and colors.

Usage: python build_deck.py <deck-spec.json> --out <deck.pptx> [--brand <brand.json>]
                             [--author "<operator name>"] [--deterministic]
Brand lookup: --brand, then the BRAND_JSON environment variable, then the nearest
brand/brand.json above this script (the project's own brand), then shared/brand.json in this
skill folder. The walk for the project's brand stops at the project root, the nearest folder at
or above this one that holds a pods/ or .claude/ folder; with no project root it checks no
folder and the shared copy is used. A "templates.pptx" path in brand.json is resolved relative
to the folder that holds brand.json; when it is absent or missing on disk the deck starts from
the python-pptx default.
--author sets the author and last-modified-by properties (default: empty). The created and
modified dates are the current time. --deterministic fixes them to STAMP and, when no
--author is given, sets the author to CREATOR, so preview and test builds are repeatable.
Titles are set at TITLE_PT and body text at BODY_PT, the SKILL.md defaults. When every bullet
on a slide starts with "1. ", "2. " and so on, the typed numbers are dropped and the slide's own
numbering is used instead, so the steps do not carry a bullet and a number.
Requires python-pptx.
"""
import argparse
import datetime
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SHARED_BRAND = os.path.join(os.path.dirname(HERE), "shared", "brand.json")
CREATOR = "ACME Co. demo generator"
STAMP = datetime.datetime(2030, 4, 1, 9, 0, 0)
LAYOUTS = {"title": 0, "bullets": 1}
TITLE_PT = 40
BODY_PT = 16
NUMBERED = re.compile(r"^\d+\.\s+")
ROOT_MARKERS = ("pods", ".claude")


def project_root(start):
    """Return the nearest folder at or above start that holds pods/ or .claude/, or None."""
    folder = os.path.abspath(start)
    while True:
        if any(os.path.isdir(os.path.join(folder, marker)) for marker in ROOT_MARKERS):
            return folder
        parent = os.path.dirname(folder)
        if parent == folder:
            return None
        folder = parent


def locate_brand(explicit=None, environ=None, start=HERE):
    environ = os.environ if environ is None else environ
    for candidate in (explicit, environ.get("BRAND_JSON")):
        if candidate:
            if not os.path.isfile(candidate):
                raise FileNotFoundError(f"brand file not found: {candidate}")
            return candidate
    root = project_root(start)
    if root is None:
        return SHARED_BRAND
    here = os.path.abspath(start)
    while True:
        candidate = os.path.join(here, "brand", "brand.json")
        if os.path.isfile(candidate):
            return candidate
        if here == root:
            return SHARED_BRAND
        here = os.path.dirname(here)


def load_spec(path):
    with open(path, encoding="utf-8") as fh:
        spec = json.load(fh)
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("spec needs a non-empty 'slides' list")
    for number, slide in enumerate(slides, 1):
        if slide.get("layout") not in LAYOUTS:
            raise ValueError(f"slide {number}: layout must be one of {sorted(LAYOUTS)}")
        if not isinstance(slide.get("title"), str) or not slide["title"]:
            raise ValueError(f"slide {number}: needs a title")
        bullets = slide.get("bullets")
        if slide["layout"] == "bullets" and (not isinstance(bullets, list) or not bullets
                                             or not all(isinstance(b, str) for b in bullets)):
            raise ValueError(f"slide {number}: a bullets slide needs a list of strings in 'bullets'")
    return spec


def _rgb(hex_value):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hex_value.lstrip("#").upper())


def _style(text_frame, face, color, size_pt):
    from pptx.util import Pt
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = face
            run.font.color.rgb = color
            run.font.size = Pt(size_pt)


def _number(paragraph):
    from pptx.oxml.ns import qn
    p_pr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buBlip"):
        for old in p_pr.findall(qn(tag)):
            p_pr.remove(old)
    auto = p_pr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"})
    following = [p_pr.find(qn(tag)) for tag in ("a:tabLst", "a:defRPr", "a:extLst")]
    anchor = next((el for el in following if el is not None), None)
    if anchor is None:
        p_pr.append(auto)
    else:
        anchor.addprevious(auto)


def build(spec, brand_path, out, author="", deterministic=False):
    from pptx import Presentation
    with open(brand_path, encoding="utf-8") as fh:
        brand = json.load(fh)
    template = (brand.get("templates") or {}).get("pptx")
    template_path = os.path.join(os.path.dirname(os.path.abspath(brand_path)), template) if template else None
    prs = Presentation(template_path) if template_path and os.path.isfile(template_path) else Presentation()
    body_face = brand["fonts"]["office"]
    heading_face = brand["fonts"].get("officeHeading", body_face)
    primary, text = _rgb(brand["colors"]["primary"]), _rgb(brand["colors"]["text"])
    for slide_spec in spec["slides"]:
        slide = prs.slides.add_slide(prs.slide_layouts[LAYOUTS[slide_spec["layout"]]])
        slide.shapes.title.text = slide_spec["title"]
        _style(slide.shapes.title.text_frame, heading_face, primary, TITLE_PT)
        body = slide.placeholders[1].text_frame
        if slide_spec["layout"] == "title":
            body.text = slide_spec.get("subtitle", "")
        else:
            bullets = slide_spec["bullets"]
            numbered = all(NUMBERED.match(b) for b in bullets)
            if numbered:
                bullets = [NUMBERED.sub("", b, count=1) for b in bullets]
            body.text = bullets[0]
            for bullet in bullets[1:]:
                body.add_paragraph().text = bullet
            if numbered:
                for paragraph in body.paragraphs:
                    _number(paragraph)
        _style(body, body_face, text, BODY_PT)
        if slide_spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = slide_spec["notes"]
    if deterministic:
        author = author or CREATOR
        stamp = STAMP
    else:
        stamp = datetime.datetime.now().replace(microsecond=0)
    props = prs.core_properties
    props.author = author
    props.last_modified_by = author
    props.title = spec.get("title", "")
    props.created = stamp
    props.modified = stamp
    props.revision = 1
    prs.save(out)
    return out


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("spec")
    parser.add_argument("--out", required=True)
    parser.add_argument("--brand", default=None)
    parser.add_argument("--author", default="",
                        help="operator name for the author and last-modified-by properties")
    parser.add_argument("--deterministic", action="store_true",
                        help="fixed dates (and CREATOR as author when --author is absent) for repeatable builds")
    args = parser.parse_args(argv)
    try:
        spec = load_spec(args.spec)
        brand = locate_brand(args.brand)
        build(spec, brand, args.out, author=args.author, deterministic=args.deterministic)
    except (OSError, ValueError, KeyError) as exc:
        print(f"build_deck: {exc}", file=sys.stderr)
        return 1
    except ImportError:
        print("build_deck: python-pptx is not installed (pip install python-pptx)", file=sys.stderr)
        return 1
    print(args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())

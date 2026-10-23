"""Build a Sankey spec (the JSON sankey.py draws) from a source table and a mapping file.

Usage:
    build_spec.py <mapping> <out_spec> [--source PATH]

The mapping says which sheets to read, which rows to keep, and how each row is placed in
every stage; references/spec-format.md lists its keys. Each kept row travels through the
stages left to right, and every adjacent pair of stages gains one unit of flow per row. A row
that no rule places stops the build, so nothing leaves the count without notice. The spec is
written only when the whole build succeeds, then a summary is printed for reconciliation.

This script reads no brand file. openpyxl is needed for workbook sources only.
"""
import os
import sys

_HERE = os.path.dirname(os.path.realpath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

import argparse  # noqa: E402
import json  # noqa: E402

from brand_palette import guard_path  # noqa: E402

PROG = "build_spec"
SHOWN_FAILURES = 20
LABEL_FALLBACK_COLUMNS = ("Name", "Title", "Summary")
WORKBOOK_TYPES = (".xlsx", ".xlsm")
OPTIONAL_COPIES = ("palette", "groups")


class BuildError(Exception):
    """A problem with the inputs; reported as one message and exit status 1."""


class Record(object):
    """One source row: the sheet it came from and its cells keyed by header."""

    __slots__ = ("sheet", "cells")

    def __init__(self, sheet, cells):
        self.sheet = sheet
        self.cells = cells

    def text(self, column):
        return as_text(self.cells.get(column))


def as_text(value):
    """Comparison form of a cell: None is blank, anything else is trimmed text."""
    return "" if value is None else str(value).strip()


# ------------------------------------------------------------------ mapping

def read_mapping(path):
    try:
        with open(path, "rb") as handle:
            raw = handle.read()
    except OSError as exc:
        raise BuildError("cannot read the mapping file %s (%s)." % (path, exc.strerror or exc))
    try:
        # json.loads refuses a leading byte-order mark, which keeps such files out.
        return json.loads(raw.decode("utf-8"))
    except ValueError as exc:
        raise BuildError("the mapping file %s is not valid JSON: %s" % (path, exc))


def check_mapping_shape(mapping):
    """Malformed rules and brackets stop the build with a message, not a traceback."""
    for stage in mapping.get("stages") or []:
        for number, rule in enumerate(stage.get("assign") or [], start=1):
            if "when" not in rule:
                raise BuildError("stage '%s', rule %d has no 'when' list. Use \"when\": [] "
                                 "for a rule that takes every row." % (stage.get("id"), number))
    for number, group in enumerate(mapping.get("groups") or [], start=1):
        if "label" not in group:
            raise BuildError("group %d has no 'label'." % number)
        if "nodes" not in group:
            raise BuildError("group %d (%r) has no 'nodes' list."
                             % (number, group.get("label")))


def source_path(mapping, mapping_file, override):
    """--source as given; else the mapping's source, joined to the mapping's folder when
    it is relative."""
    if override:
        return override
    named = mapping.get("source")
    if not named:
        raise BuildError("the mapping has no 'source' and --source was not given.")
    # Checked before the join, so a /d/... value is refused on every Python version.
    guard_path(named, PROG)
    if os.path.isabs(named):
        return named
    mapping_folder = os.path.dirname(os.path.abspath(mapping_file))
    return os.path.join(mapping_folder, named)


# ------------------------------------------------------------------ source table

def require_columns(where, headers, entry):
    """A sheet's own key column and filter columns must be in its header."""
    wanted = set()
    if entry.get("key_column"):
        wanted.add(entry["key_column"])
    for rule in entry.get("filters") or []:
        wanted.add(rule.get("column"))
    for column in sorted(wanted, key=str):
        if column not in headers:
            raise BuildError("%s has no column '%s'. Columns found: %s"
                             % (where, column, ", ".join(str(h) for h in headers)))


def is_empty_header(value):
    return value is None or value == ""


def read_workbook(path, sheets):
    import openpyxl  # only workbook sources need it

    book = openpyxl.load_workbook(path, data_only=True)
    records = []
    seen = set()
    for entry in sheets:
        name = entry.get("name")
        if name not in book.sheetnames:
            raise BuildError("sheet '%s' is not in %s. Sheets present: %s"
                             % (name, path, ", ".join(book.sheetnames)))
        rows = book[name].iter_rows(values_only=True)
        header = list(next(rows, ()))
        named = [h for h in header if not is_empty_header(h)]
        require_columns("sheet '%s'" % name, named, entry)
        seen.update(named)
        key = entry.get("key_column")
        for row in rows:
            pairs = list(zip(header, row))
            cells = dict(pairs)
            if key:
                if as_text(cells.get(key)) == "":
                    continue
            elif all(as_text(value) == "" for _, value in pairs):
                continue
            records.append(Record(name, cells))
    return records, seen


def read_csv(path, sheets):
    import csv

    first = sheets[0] if sheets else {}
    if sheets and "name" not in first:
        raise BuildError("the first 'sheets' entry has no 'name'; a CSV uses it as the "
                         "sheet name for its rows.")
    tag = first.get("name", "")
    key = first.get("key_column")
    records = []
    with open(path, encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        header = next(reader, [])
        named = [h for h in header if h != ""]
        require_columns(os.path.basename(path), named, first)
        for row in reader:
            if not row:
                continue
            cells = {}
            for pos, column in enumerate(header):
                cells[column] = row[pos] if pos < len(row) else ""
            if key and as_text(cells.get(key)) == "":
                continue
            records.append(Record(tag, cells))
    return records, set(named)


def read_source(path, mapping):
    guard_path(path, PROG)
    if not os.path.exists(path):
        raise BuildError("source file not found: %s" % path)
    kind = os.path.splitext(path)[1].lower()
    if kind not in WORKBOOK_TYPES and kind != ".csv":
        raise BuildError("source type '%s' is not supported; give a .xlsx or .csv file."
                         % kind)
    if "sheets" not in mapping:
        raise BuildError("the mapping has no 'sheets' list. Give one entry per sheet, or "
                         "\"sheets\": [] for a CSV with no key column.")
    sheets = mapping["sheets"]
    if kind == ".csv":
        return read_csv(path, sheets)
    return read_workbook(path, sheets)


def check_rule_columns(stages, seen):
    """A rule column may be absent from some sheets, but not from all of them."""
    named = set()
    for stage in stages:
        for rule in stage.get("assign") or []:
            for cond in rule.get("when") or []:
                if "column" in cond:
                    named.add(cond["column"])
    unknown = sorted((c for c in named if c not in seen), key=str)
    if unknown:
        raise BuildError("assignment rules use column(s) found on no loaded sheet: %s. "
                         "Check for a typo." % ", ".join(str(c) for c in unknown))


# ------------------------------------------------------------------ rule tests

_MEMBERSHIP = (
    ("not_in", lambda value, listed: value not in listed),
    ("in", lambda value, listed: value in listed),
    ("equals", lambda value, wanted: value == wanted),
)


def membership_holds(value, test):
    """not_in, in and equals: each one present in the test must hold for value."""
    return all(check(value, test[key]) for key, check in _MEMBERSHIP if key in test)


def value_passes(value, test):
    """Filter or column condition on one normalized cell."""
    is_blank = not value
    if test.get("blank") and not is_blank:
        return False
    if test.get("not_blank") and is_blank:
        return False
    return membership_holds(value, test)


def condition_holds(cond, record, placed):
    if "column" in cond:
        if not value_passes(record.text(cond["column"]), cond):
            return False
    if "stage" in cond:
        # An unknown or later stage reads as None here.
        if not membership_holds(placed.get(cond["stage"]), cond):
            return False
    if "sheet" in cond:
        if not cond["sheet"] == record.sheet:
            return False
    return True


def sheet_entry(sheets, tag):
    for entry in sheets:
        if entry.get("name", "") == tag:
            return entry
    return None


def keep_record(record, sheets):
    entry = sheet_entry(sheets, record.sheet)
    rules = (entry or {}).get("filters") or []
    return all(value_passes(record.text(rule.get("column")), rule) for rule in rules)


def describe(record, sheets, label_columns):
    entry = sheet_entry(sheets, record.sheet) or {}
    candidates = []
    if entry.get("key_column"):
        candidates.append(entry["key_column"])
    candidates.extend(label_columns)
    candidates.extend(LABEL_FALLBACK_COLUMNS)
    for column in candidates:
        shown = record.text(column)
        if shown:
            return shown
    return "<unnamed>"


def miss_lines(misses, sheets, label_columns):
    """The documented entry form, one per unplaced row, capped at SHOWN_FAILURES."""
    for record, stage_id in misses[:SHOWN_FAILURES]:
        yield "  - %s (sheet %s) matched no rule in stage '%s'" % (
            describe(record, sheets, label_columns), record.sheet, stage_id)


# ------------------------------------------------------------------ the build

def stage_nodes(stage):
    """Declared (node id, label, color) triples; labels and colors pass through as given
    and default only when the key is absent."""
    fixed = stage.get("constant")
    if fixed:
        node = fixed["node"]
        return [(node, fixed.get("label", node), fixed.get("color", "total"))]
    return [(n["id"], n.get("label", n["id"]), n.get("color", "neutral"))
            for n in stage.get("nodes") or []]


def place(record, stages):
    """Node ids for the record, one per stage, or (None, failing stage id)."""
    placed = {}
    route = []
    for stage in stages:
        fixed = stage.get("constant")
        if fixed:
            node = fixed["node"]
        else:
            for rule in stage.get("assign") or []:
                if all(condition_holds(c, record, placed) for c in rule["when"]):
                    node = rule["node"]
                    break
            else:
                return None, stage["id"]
        placed[stage["id"]] = node
        route.append(node)
    return route, None


def build(mapping, records):
    """Return (spec, notes). Raises BuildError on any failure."""
    sheets = mapping["sheets"]
    stages = mapping["stages"]
    stage_ids = [stage["id"] for stage in stages]

    kept = [r for r in records if keep_record(r, sheets)]
    if not kept:
        raise BuildError("every row was filtered out; confirm the sheet names and filter "
                         "values in the mapping.")

    routes = []
    misses = []
    for record in kept:
        route, failed_stage = place(record, stages)
        if route is None:
            misses.append((record, failed_stage))
        else:
            routes.append(route)
    if misses:
        head = ("%d row(s) matched no assignment rule. Rows are never dropped quietly; add "
                "or fix a rule so each one lands in a node." % len(misses))
        listed = miss_lines(misses, sheets, mapping.get("label_columns") or [])
        raise BuildError("\n".join([head] + list(listed)))

    per_stage = dict((sid, {}) for sid in stage_ids)
    flows = {}
    for route in routes:
        for sid, node in zip(stage_ids, route):
            tally = per_stage[sid]
            tally[node] = tally.get(node, 0) + 1
        for pair in zip(route, route[1:]):
            flows[pair] = flows.get(pair, 0) + 1

    carrying = set()
    for left, right in flows:
        carrying.add(left)
        carrying.add(right)

    columns = []
    notes = []
    for pos, stage in enumerate(stages):
        declared = [node for node, _, _ in stage_nodes(stage)]
        present = [node for node in declared if node in carrying]
        if not present:
            raise BuildError("stage %d ('%s') has no nodes carrying data." % (pos, stage["id"]))
        for node in declared:
            if node not in carrying:
                notes.append("NOTE: node '%s' matched no rows; it is left off the chart."
                             % node)
        columns.append(present)

    labels = {}
    colors = {}
    for stage in stages:
        for node, label, color in stage_nodes(stage):
            labels[node] = label
            colors[node] = color

    endpoint = dict(per_stage[stage_ids[-1]])
    groups = mapping.get("groups") or []
    default_sides = ["right"] + ["left"] * (len(stages) - 1)

    spec = {}
    for key in ("title", "subtitle", "footer", "prepared_by"):
        spec[key] = mapping.get(key, "")
    spec["column_titles"] = [stage.get("title", stage["id"]) for stage in stages]
    spec["columns"] = columns
    spec["label_sides"] = mapping.get("label_sides") or default_sides
    spec["labels"] = labels
    spec["colors"] = colors
    spec["links"] = [{"source": s, "target": t, "value": n} for (s, t), n in flows.items()]
    spec["color_by"] = mapping.get("color_by", "target")
    for key in OPTIONAL_COPIES:
        if mapping.get(key):
            spec[key] = mapping[key]
    spec["_population"] = len(routes)
    spec["_stage_counts"] = per_stage
    spec["_endpoint_counts"] = endpoint
    if groups:
        spec["_group_counts"] = dict((g["label"], group_total(g, endpoint)) for g in groups)
    return spec, notes


def group_total(group, endpoint):
    return sum(endpoint.get(node, 0) for node in group["nodes"])


def write_spec(path, spec):
    try:
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(json.dumps(spec, indent=1))
    except OSError as exc:
        raise BuildError("cannot write %s (%s)." % (path, exc.strerror or exc))


def print_summary(spec, out_path):
    print("population: %d" % spec["_population"])
    endpoint = spec["_endpoint_counts"]
    # sorted() stays stable with reverse=True, so tied counts keep first-seen order.
    for node, count in sorted(endpoint.items(), key=lambda kv: kv[1], reverse=True):
        print("  %-28s %d" % (node, count))
    for label, count in spec.get("_group_counts", {}).items():
        print("  [group] %-20s %d" % (label, count))
    print(out_path)


# ------------------------------------------------------------------ command line

def parse_args(argv):
    parser = argparse.ArgumentParser(
        prog="build_spec.py",
        usage="%(prog)s <mapping> <out_spec> [--source PATH]",
        description="Build a Sankey spec JSON from a source table and a mapping file.")
    parser.add_argument("mapping", help="the mapping file: sheets, filters, stages, rules")
    parser.add_argument("out_spec", help="where the finished spec is saved")
    parser.add_argument("--source", default=None,
                        help="read this table in place of the mapping's own source; a "
                             "relative path counts from the folder you run in")
    return parser.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    guard_path(args.mapping, PROG)
    guard_path(args.out_spec, PROG)
    try:
        mapping = read_mapping(args.mapping)
        check_mapping_shape(mapping)
        table = source_path(mapping, args.mapping, args.source)
        records, seen = read_source(table, mapping)
        check_rule_columns(mapping["stages"], seen)
        spec, notes = build(mapping, records)
        for note in notes:
            print(note, file=sys.stderr)
        write_spec(args.out_spec, spec)
    except BuildError as exc:
        print("%s: %s" % (PROG, exc), file=sys.stderr)
        return 1
    print_summary(spec, args.out_spec)
    return 0


if __name__ == "__main__":
    sys.exit(main())

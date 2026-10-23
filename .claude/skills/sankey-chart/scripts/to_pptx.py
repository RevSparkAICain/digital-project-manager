"""Put a rendered Sankey image on a PowerPoint slide.

    python scripts/to_pptx.py <png> <out_pptx> [options]

The base deck is, in order of preference, the --template file, the template the brand file
names, or a blank widescreen deck. A template is copied to the output path before anything
is added, so the original stays as it was and its slides are kept. One slide is appended,
cleared of placeholders, and given the optional title and takeaway plus the image, scaled to
fit inside half-inch margins with its proportions kept.
"""
import argparse
import os
import re
import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from brand_palette import guard_path, load_brand, lookup  # noqa: E402

PROG = "to_pptx"

MARGIN = Inches(0.5)
TITLE_BOX_H = Inches(0.6)
TITLE_STEP = Inches(0.7)
TAKEAWAY_BOX_H = Inches(0.45)
TAKEAWAY_STEP = Inches(0.55)
BLANK_WIDTH = Inches(13.333)
BLANK_HEIGHT = Inches(7.5)

_HEX6 = re.compile(r"#?([0-9A-Fa-f]{6})")


def brand_rgb(brand, key):
    raw = lookup(brand, key)
    hit = _HEX6.fullmatch(raw) if isinstance(raw, str) else None
    if hit is None:
        sys.exit("%s: brand value %s must be a six-digit hex color, found %r"
                 % (PROG, key, raw))
    return RGBColor.from_string(hit.group(1).upper())


def brand_template(brand, brand_path):
    """Template path the brand names, or None. A named file that is absent gets a note."""
    named = lookup(brand, "templates.pptx")
    if not named:
        return None
    if os.path.isabs(named):
        where = named
    else:
        where = str(Path(os.path.abspath(brand_path)).parent / named)
    if os.path.exists(where):
        return where
    print("NOTE: the brand template '%s' is missing (looked for %s); a blank deck is used "
          "instead." % (named, where), file=sys.stderr)
    return None


def open_deck(template, out_pptx):
    if template:
        template = guard_path(template, PROG)
        if not os.path.exists(template):
            sys.exit("%s: template not found: %s" % (PROG, template))
        shutil.copyfile(template, out_pptx)
        return Presentation(out_pptx)
    deck = Presentation()
    deck.slide_width = BLANK_WIDTH
    deck.slide_height = BLANK_HEIGHT
    return deck


def choose_layout(deck, wanted):
    layouts = list(deck.slide_layouts)
    if wanted.strip():
        key = wanted.strip().lower()
        for layout in layouts:
            if layout.name.strip().lower() == key:
                return layout
        print("NOTE: layout '%s' is not in this deck (layouts: %s); the layout with the "
              "fewest placeholders is used instead."
              % (wanted, ", ".join(layout.name for layout in layouts)), file=sys.stderr)
    # min keeps the first of equals, so ties go to master order.
    return min(layouts, key=lambda layout: len(layout.placeholders))


def add_line(slide, top, height, width, text, look):
    """One wrapped text box with a single run; look is (points, bold, rgb, face)."""
    points, bold, rgb, face = look
    box = slide.shapes.add_textbox(MARGIN, top, width, height)
    box.text_frame.word_wrap = True
    run = box.text_frame.paragraphs[0].add_run()
    font = run.font
    if face:
        font.name = face
    font.color.rgb = rgb
    font.bold = bold
    font.size = Pt(points)
    run.text = text


def build(args):
    from PIL import Image

    png = guard_path(args.png, PROG)
    if not os.path.exists(png):
        sys.exit("%s: image not found: %s" % (PROG, png))
    out_pptx = guard_path(args.out_pptx, PROG)

    brand, brand_path = load_brand(args.brand or None)
    title_color = brand_rgb(brand, "colors.text")
    takeaway_color = brand_rgb(brand, "colors.textMuted")
    body_face = lookup(brand, "fonts.office") or ""
    heading_face = lookup(brand, "fonts.officeHeading") or body_face

    template = args.template or brand_template(brand, brand_path)
    deck = open_deck(template, out_pptx)

    slide = deck.slides.add_slide(choose_layout(deck, args.layout))
    for holder in list(slide.placeholders):
        node = holder.element
        node.getparent().remove(node)

    slide_w, slide_h = deck.slide_width, deck.slide_height
    inner_w = slide_w - 2 * MARGIN
    top = MARGIN
    if args.title:
        add_line(slide, top, TITLE_BOX_H, inner_w, args.title,
                 (28, True, title_color, heading_face))
        top += TITLE_STEP
    if args.takeaway:
        add_line(slide, top, TAKEAWAY_BOX_H, inner_w, args.takeaway,
                 (14, False, takeaway_color, body_face))
        top += TAKEAWAY_STEP

    with Image.open(png) as picture:
        px_w, px_h = picture.size
    room_h = slide_h - top - MARGIN
    ratio = min(inner_w / float(px_w), room_h / float(px_h))
    pic_w, pic_h = int(px_w * ratio), int(px_h * ratio)
    left = int((slide_w - pic_w) / 2.0)
    pic_top = int(top + (room_h - pic_h) / 2.0)
    slide.shapes.add_picture(png, left, pic_top, width=pic_w, height=pic_h)

    if args.author:
        deck.core_properties.author = args.author
        deck.core_properties.last_modified_by = args.author

    deck.save(out_pptx)


def parse_args(argv):
    parser = argparse.ArgumentParser(
        prog="to_pptx.py", description="Append one slide holding the chart image to a deck.")
    parser.add_argument("png", help="chart image to place")
    parser.add_argument("out_pptx", help="deck to write")
    parser.add_argument("--template", default="", help="base deck to copy and add to")
    parser.add_argument("--title", default="", help="slide title")
    parser.add_argument("--takeaway", default="", help="one line under the title")
    parser.add_argument("--author", default="", help="author and last-modified-by name")
    parser.add_argument("--layout", default="", help="slide layout name, any case")
    parser.add_argument("--brand", default="", help="brand JSON file")
    return parser.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    build(args)
    print(args.out_pptx)


if __name__ == "__main__":
    main()

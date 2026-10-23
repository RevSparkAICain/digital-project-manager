"""Generate the synthetic ACME Co. support-desk workbook used by the sankey-chart demo.

Usage:
    python make_demo_data.py [out.xlsx]        (default: tickets.xlsx beside this file)

Every value is invented. The output is deterministic: the same counts, the same ticket
keys and the same row order on every run, so expected-counts.json stays valid.

Sheets:
    Service Desk     342 rows: 340 tickets plus 2 note rows with no Ticket key
    Field Requests    80 rows: 60 in scope (Y), 15 out of scope (N), 5 with In Scope blank
    Escalations       25 rows: copies of Service Desk tickets under their own keys
"""
import datetime
import os
import random
import sys

# (type, endpoint bucket) -> count. Rows per type: Access 120, Software 100,
# Hardware 70, Network 50. Field Requests in scope: 60.
SERVICE_DESK = {
    "Access":   {"done": 70, "duplicate": 10, "canceled": 5, "waiting": 10, "working": 15, "new": 10},
    "Software": {"done": 50, "duplicate": 10, "canceled": 5, "waiting": 10, "working": 15, "new": 10},
    "Hardware": {"done": 30, "duplicate": 0,  "canceled": 5, "waiting": 5,  "working": 20, "new": 10},
    "Network":  {"done": 30, "duplicate": 0,  "canceled": 0, "waiting": 5,  "working": 10, "new": 5},
}
FIELD_IN_SCOPE = {"done": 20, "canceled": 5, "waiting": 10, "working": 15, "new": 10}
FIELD_OUT_OF_SCOPE = 15
FIELD_BLANK_SCOPE = 5
ESCALATIONS = 25

SUMMARIES = {
    "Access": ["Password reset", "New hire account", "Group membership change",
               "MFA device swap", "Shared mailbox access"],
    "Software": ["Install request", "License renewal", "App crashes on launch",
                 "Update failed", "Add-in missing"],
    "Hardware": ["Laptop will not boot", "Dock not detected", "Replace keyboard",
                 "Monitor flicker", "Printer jam"],
    "Network": ["Wi-Fi drops", "VPN slow", "Port activation", "DNS lookup fails",
                "Guest network request"],
    "Field": ["Site visit", "Conference room setup", "Cable run", "Badge reader check",
              "Kiosk swap"],
}
SITES = ["North campus", "South campus", "Warehouse", "Remote"]


def status_for(bucket, i, has_resolution):
    """Map an endpoint bucket to the Status (and Resolution) a real desk would record."""
    if bucket == "done":
        return ("Closed" if i % 2 == 0 else "Resolved"), ("Fixed" if has_resolution else None)
    if bucket == "duplicate":
        # Closed like finished work. Only the Resolution column says otherwise, which is
        # why the mapping's Duplicate rule must run before the Closed rule.
        return "Closed", "Duplicate"
    if bucket == "canceled":
        return "Canceled", None
    if bucket == "waiting":
        return "Waiting on Customer", None
    if bucket == "working":
        return "In Progress", None
    return "New", None


def opened(rng):
    return datetime.date(2030, 1, 1) + datetime.timedelta(days=rng.randrange(0, 90))


def build_rows(rng):
    desk = []
    for typ, buckets in SERVICE_DESK.items():
        for bucket, n in buckets.items():
            for i in range(n):
                status, res = status_for(bucket, i, True)
                desk.append({"Summary": rng.choice(SUMMARIES[typ]), "Type": typ,
                             "Status": status, "Resolution": res, "Opened": opened(rng)})
    rng.shuffle(desk)

    field = []
    for bucket, n in FIELD_IN_SCOPE.items():
        for i in range(n):
            status, _ = status_for(bucket, i, False)
            field.append({"Summary": rng.choice(SUMMARIES["Field"]), "Site": rng.choice(SITES),
                          "Status": status, "In Scope": "Y", "Opened": opened(rng)})
    for i in range(FIELD_OUT_OF_SCOPE):
        field.append({"Summary": rng.choice(SUMMARIES["Field"]), "Site": rng.choice(SITES),
                      "Status": ["New", "Closed", "Canceled"][i % 3], "In Scope": "N",
                      "Opened": opened(rng)})
    for i in range(FIELD_BLANK_SCOPE):
        field.append({"Summary": rng.choice(SUMMARIES["Field"]), "Site": rng.choice(SITES),
                      "Status": "New", "In Scope": None, "Opened": opened(rng)})
    rng.shuffle(field)

    key = 101
    for r in desk:
        r["Ticket"] = "DEMO-%d" % key
        key += 1
    key = 501
    for r in field:
        r["Ticket"] = "DEMO-%d" % key
        key += 1

    parents = rng.sample(desk, ESCALATIONS)
    esc = []
    for i, p in enumerate(parents):
        esc.append({"Ticket": "DEMO-%d" % (701 + i), "Parent": p["Ticket"],
                    "Summary": p["Summary"], "Type": p["Type"], "Status": p["Status"]})
    return desk, field, esc


def write(path):
    import openpyxl
    rng = random.Random(2030)
    desk, field, esc = build_rows(rng)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Service Desk"
    cols = ["Ticket", "Summary", "Type", "Status", "Resolution", "Opened"]
    ws.append(cols)
    for i, r in enumerate(desk):
        ws.append([r.get(c) for c in cols])
        if i == 99:     # a note row with no Ticket key: build_spec skips it via key_column
            ws.append([None, "Note: rows below imported from the old queue", None, None,
                       None, None])
    ws.append([None, "Note: end of export", None, None, None, None])

    ws = wb.create_sheet("Field Requests")
    cols = ["Ticket", "Summary", "Site", "Status", "In Scope", "Opened"]
    ws.append(cols)
    for r in field:
        ws.append([r.get(c) for c in cols])

    ws = wb.create_sheet("Escalations")
    cols = ["Ticket", "Parent", "Summary", "Type", "Status"]
    ws.append(cols)
    for r in esc:
        ws.append([r.get(c) for c in cols])

    stamp = datetime.datetime(2030, 4, 1, 9, 0, 0)
    wb.properties.creator = "ACME Co. demo generator"
    wb.properties.created = stamp
    wb.properties.modified = stamp
    wb.save(path)
    pin_timestamps(path, stamp)
    return path


def pin_timestamps(path, stamp):
    """openpyxl stamps 'modified' and every zip entry with the current time on save.

    Rewrite both to the fixed demo date so the workbook carries no real date and the
    file is byte-identical across runs.
    """
    import re
    import zipfile
    with zipfile.ZipFile(path) as zin:
        entries = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]
    iso = stamp.strftime("%Y-%m-%dT%H:%M:%SZ").encode("ascii")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries:
            if name == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + iso + rb"\g<2>", data)
            info = zipfile.ZipInfo(name, date_time=stamp.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            zout.writestr(info, data)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "tickets.xlsx")
    print(write(out))

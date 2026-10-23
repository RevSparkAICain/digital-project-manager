"""Render executive meeting notes written in Markdown as a branded Word document.

Usage:
    python md2docx.py <notes.md> --out <file.docx> --title "<meeting name>"
        --topic "<what the meeting covered>" --date "<Month D, YYYY>"
        --operator "<person running the session>" --doc-type "<category>"
        --type-label "<page label>" [--footer "<left footer text>"] [--brand <brand.json>]

Exit codes: 0 written; 2 an argument was refused (nothing is written); 1 anything else that
stopped the build (missing input or brand file, python-docx not installed, a theme slot the
retarget step could not find). Every failure prints one line on stderr.

Brand lookup, first hit wins:
    1. --brand
    2. the BRAND_JSON environment variable
    3. the nearest brand/brand.json in this folder or a folder above it, up to and including
       the project root (the project's own)
    4. shared/brand.json in this skill folder
A path named by 1 or 2 that does not exist is an error, never a silent fall-through. The
project root is the nearest folder at or above this one that holds a pods/ or .claude/
folder; the walk in 3 stops there. With no project root, 3 checks no folder.

Refusals, checked before the output path is touched:
    --operator    blank, wrapped in angle brackets, or a stand-in such as "Operator Name",
                  "your name", "unknown" or "TBD". The operator is stamped as author and in the
                  byline; a wrong name is worse than no document, so the caller must pass it.
    --title       blank. The title paragraph and the default footer are built from it.
    --topic       blank, holding a pipe or "prepared by", or equal to --doc-type or
                  --type-label. The script composes the byline itself, so a topic that already
                  carries a byline or a type would print those parts twice.
    --doc-type    blank, or equal to --type-label. The two flags carry different words on
    --type-label  purpose (a category for the properties, a label for the page); neither has a
                  default, so a swap or a copy cannot slip through unseen.
    --date        not "Month D, YYYY" with a full English month name and a real calendar day.
                  Month names come from the MONTHS table, not the machine locale, and the date
                  is printed back in one canonical form ("March 3, 2030").
    --out         the notes file itself, however the path is spelled. Saving there would
                  replace the Markdown with the document.
Comparisons trim, collapse inner whitespace and ignore case.

What the document gets:
    - Title paragraph, then a byline: "<topic> | <date> | <type label> | Prepared by <operator>".
    - Body from a small Markdown subset: "#", "##", "###" headings; "- " bullets; "N. " numbered
      items (one to three digits); "**bold**" spans; anything else is a plain paragraph. A
      leading label of up to four words ending in a colon ("Decision:", "Sign-off:",
      "Budget (draft):") is set bold on bullets and numbered items. A line that is only a bold
      label keeps with the next paragraph. A blank line, heading, bullet or paragraph ends a
      numbered list, and every list restarts at 1 on its own numbering instance.
    - Title block: a bar in colors.primary runs down the left of the title and the byline,
      and the byline is bold in colors.primary.
    - First-page header: a band across the top edge, colors.primary into colors.accent1 at
      62/38, then logo.lockup right-aligned. Later pages carry no header. A brand without
      logo.lockup gets no band and no logos; one that names a missing file stops the build.
    - Footer, every page: left text (default "<title>, <date>"), a right tab, the lockup (when
      the brand has one), then "Page X of Y" as live fields, bold in colors.primary. The
      paragraph clears the tab stops the Footer style would add.
    - Bullets are small squares in colors.primary. Heading 1 has a rule under it in
      colors.divider.
    - Core properties: author and last modified by = operator; title = "<title> - <doc type>";
      subject = footer left text; category = doc type; comments = "Prepared by <operator>";
      created and modified = the meeting date at midnight; revision 1.
    - Styles: body and list styles in fonts.office and colors.text; Title and Heading 1 in
      fonts.officeHeading and colors.heading (colors.primary when the brand has no heading
      colour) at large-text size; Heading 2 and 3 below the large-text line, so they take the
      first of primary, link, text reaching 4.5:1 on colors.background (WCAG 2.x). No style
      colour is left bound to the theme.
    - Theme: colour slots and the two theme fonts are rewritten from the brand file, mapped as
      accent1 primary, accent2 text, accent3 textMuted, accent4 link, accent5 divider,
      accent6 surface, dk2 text, lt2 surface, hlink link, folHlink link; major font
      fonts.officeHeading, minor font fonts.office. Each slot must appear exactly once.

Standard library for argument checks; python-docx (imported inside the build functions) for
the document itself.
"""
import argparse
import datetime
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SHARED_BRAND = os.path.join(os.path.dirname(HERE), "shared", "brand.json")
ROOT_MARKERS = ("pods", ".claude")

MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August",
          "September", "October", "November", "December")
_MONTH_NUMBER = {name.casefold(): number for number, name in enumerate(MONTHS, 1)}
_DATE = re.compile(r"(?P<month>[A-Za-z]+) +(?P<day>[0-9][0-9]?), *(?P<year>[0-9]{4})")

_WML = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _w(*tags):
    """Clark-notation name, or a slash-joined path of names, in the WordprocessingML namespace."""
    return "/".join(f"{{{_WML}}}{tag}" for tag in tags)


STAND_IN_OPERATORS = {"operator name", "operator", "name", "your name", "unknown", "tbd", "n/a",
                      "python-docx"}

LARGE_TEXT_PT = 14
SMALL_TEXT_RATIO = 4.5
STYLE_SIZES_PT = {"Normal": 11, "Title": 26, "Heading 1": 16, "Heading 2": 13, "Heading 3": 12}
BYLINE_PT = 10
FOOTER_PT = 9
BAND_IN = 0.22                # first-page band height
BAND_SPLIT = 0.62             # primary share of the band; accent1 takes the rest
LOCKUP_HEADER_IN = 0.55       # logo heights, first-page header and footer
LOCKUP_FOOTER_IN = 0.22
TITLE_BAR = (48, 12)          # left bar on the title block: width in eighths of a point, gap in points
H1_RULE = (6, 4)              # rule under Heading 1, same units
BULLET = "▪"             # small square, in fonts.office
_PBDR_SUCCESSORS = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
                    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                    "w:pPrChange")
LIST_STYLES = ("List Bullet", "List Bullet 2", "List Number", "List Number 2")

THEME_COLOR_SLOTS = (("accent1", "primary"), ("accent2", "text"), ("accent3", "textMuted"),
                     ("accent4", "link"), ("accent5", "divider"), ("accent6", "surface"),
                     ("dk2", "text"), ("lt2", "surface"), ("hlink", "link"),
                     ("folHlink", "link"))
THEME_FONT_GROUPS = (("majorFont", "officeHeading"), ("minorFont", "office"))
_THEME_ATTRS = ("themeColor", "themeShade", "themeTint")
_THEME_FONT_ATTRS = ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")

_HEADING = re.compile(r"(#{1,3})\s+(.*\S)")
_BULLET = re.compile(r"-\s+(.*\S)")
_NUMBERED = re.compile(r"\d{1,3}\.\s+(.*\S)")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_BOLD_LABEL_LINE = re.compile(r"\*\*([^*]+:)\*\*|\*\*([^*]+)\*\*:")
_WORD = r"[A-Za-z]+(?:[-/][A-Za-z]+)*"
_LABEL = re.compile(rf"{_WORD}(?: {_WORD}){{0,3}}(?: \([A-Za-z]+(?: [A-Za-z]+){{0,2}}\))?:(?=\s|$)")


# ---------------------------------------------------------------- brand and colour


def project_root(start):
    """Return the nearest folder at or above start that holds pods/ or .claude/, or None."""
    folder = os.path.abspath(start)
    while True:
        if any(os.path.isdir(os.path.join(folder, marker)) for marker in ROOT_MARKERS):
            return folder
        parent = os.path.dirname(folder)
        if parent == folder:
            return None
        folder = parent


def locate_brand(explicit=None, environ=None, start=HERE):
    """Return the brand file path by the lookup order in the module docstring."""
    environ = os.environ if environ is None else environ
    for candidate in (explicit, environ.get("BRAND_JSON")):
        if candidate:
            if not os.path.isfile(candidate):
                raise FileNotFoundError(f"brand file not found: {candidate}")
            return candidate
    root = project_root(start)
    if root is None:
        return SHARED_BRAND
    folder = os.path.abspath(start)
    while True:
        candidate = os.path.join(folder, "brand", "brand.json")
        if os.path.isfile(candidate):
            return candidate
        if folder == root:
            return SHARED_BRAND
        folder = os.path.dirname(folder)


def _hex(value):
    digits = str(value).strip().lstrip("#").upper()
    if not re.fullmatch(r"[0-9A-F]{6}", digits):
        raise ValueError(f"not a six-digit hex colour: {value!r}")
    return digits


def _luminance(value):
    digits = _hex(value)
    linear = []
    for start in (0, 2, 4):
        channel = int(digits[start:start + 2], 16) / 255
        linear.append(channel / 12.92 if channel <= 0.04045
                      else ((channel + 0.055) / 1.055) ** 2.4)
    red, green, blue = linear
    return 0.2126 * red + 0.7152 * green + 0.0722 * blue


def contrast(fg_hex, bg_hex):
    """WCAG 2.x contrast ratio between two hex colours, from 1.0 to 21.0."""
    lighter, darker = sorted((_luminance(fg_hex), _luminance(bg_hex)), reverse=True)
    return (lighter + 0.05) / (darker + 0.05)


def small_heading_color(colors):
    """First of primary, link, text that reaches 4.5:1 on the background; text if none does."""
    for key in ("primary", "link", "text"):
        value = colors.get(key)
        if value and contrast(value, colors["background"]) >= SMALL_TEXT_RATIO:
            return value
    return colors["text"]


# ---------------------------------------------------------------- argument checks


def _plain(value):
    return " ".join(str(value).split()).casefold()


def _refuse(flag, problem, wanted):
    raise ValueError(f"{flag} {problem}; it should hold {wanted}")


def _parse_date(text):
    match = _DATE.fullmatch(str(text).strip())
    if not match:
        _refuse("--date", f"is {text!r}, not in Month D, YYYY form",
                'a date such as "March 12, 2030"')
    month = _MONTH_NUMBER.get(match["month"].casefold())
    if month is None:
        _refuse("--date", f"names {match['month']!r}, which is not an English month",
                "a full English month name, day and year")
    try:
        return datetime.datetime(int(match["year"]), month, int(match["day"]))
    except ValueError:
        _refuse("--date", f"is {text!r}, which is not a real calendar day",
                "a date that exists")


def format_date(moment):
    """Canonical page form of a date: full English month, unpadded day, four-digit year."""
    return f"{MONTHS[moment.month - 1]} {moment.day}, {moment.year}"


def validate(*, title, topic, date, operator, doc_type, type_label):
    """Check every argument; raise ValueError naming the flag, or return the meeting date."""
    who = _plain(operator)
    raw = str(operator).strip()
    if not who:
        _refuse("--operator", "is blank", "the name of the person running the session")
    if (raw.startswith("<") and raw.endswith(">")) or who in STAND_IN_OPERATORS:
        _refuse("--operator", f"is a placeholder ({raw!r})",
                "the real name of the person running the session")
    if not _plain(title):
        _refuse("--title", "is blank", "the meeting name")
    subject = _plain(topic)
    if not subject:
        _refuse("--topic", "is blank", "what the meeting was about")
    if "|" in subject or "prepared by" in subject:
        _refuse("--topic", f"looks like a finished byline ({str(topic).strip()!r})",
                "only what the meeting was about; the script builds the byline")
    if not _plain(doc_type):
        _refuse("--doc-type", "is blank", "the document category from the config")
    if not _plain(type_label):
        _refuse("--type-label", "is blank", "the page label from the config")
    if subject == _plain(doc_type):
        _refuse("--topic", "repeats --doc-type", "what the meeting was about, not its type")
    if subject == _plain(type_label):
        _refuse("--topic", "repeats --type-label", "what the meeting was about, not its label")
    if _plain(doc_type) == _plain(type_label):
        _refuse("--doc-type", "matches --type-label (were the two swapped or copied?)",
                "the category, with --type-label holding the different page label")
    return _parse_date(date)


# ---------------------------------------------------------------- theme


def retarget_theme(xml, colors, fonts):
    """Point the theme colour slots and theme fonts at brand values; return the new XML text.

    Raises KeyError when a slot or font group is not present exactly once, so a changed
    template stops the build instead of shipping with stock values left in place.
    """
    from xml.sax.saxutils import quoteattr

    for slot, key in THEME_COLOR_SLOTS:
        pattern = re.compile(rf"(<a:{slot}>).*?(</a:{slot}>)", re.S)
        hits = len(pattern.findall(xml))
        if hits != 1:
            raise KeyError(f"theme colour slot {slot} found {hits} times, expected once")
        value = _hex(colors[key])
        xml = pattern.sub(lambda m: f'{m.group(1)}<a:srgbClr val="{value}"/>{m.group(2)}', xml)
    for group, key in THEME_FONT_GROUPS:
        pattern = re.compile(rf'(<a:{group}>\s*<a:latin\s+typeface=)"[^"]*"')
        hits = len(pattern.findall(xml))
        if hits != 1:
            raise KeyError(f"theme font group {group} found {hits} times, expected once")
        face = quoteattr(fonts[key])
        xml = pattern.sub(lambda m: m.group(1) + face, xml)
    return xml


def _part(document, suffix):
    for part in document.part.package.iter_parts():
        if str(part.partname).endswith(suffix):
            return part
    return None


def _apply_theme(document, colors, fonts):
    theme = _part(document, "/theme/theme1.xml")
    if theme is None:
        raise KeyError("theme part /word/theme/theme1.xml not found")
    text = retarget_theme(theme.blob.decode("utf-8"), colors, fonts)
    theme._blob = text.encode("utf-8")


# ---------------------------------------------------------------- styles


def _drop_attrs(element, names):
    for name in names:
        element.attrib.pop(_w(name), None)


def _style_font(font_holder, element, face, colour, size_pt=None, bold=None):
    """Set face, colour, size and weight on a style or run, leaving nothing theme-bound."""
    from docx.shared import Pt, RGBColor

    font = font_holder.font
    font.name = face
    font.color.rgb = RGBColor.from_string(_hex(colour))
    if size_pt is not None:
        font.size = Pt(size_pt)
    if bold is not None:
        font.bold = bold
    rpr = element.find(_w("rPr"))
    faces = rpr.find(_w("rFonts"))
    _drop_attrs(faces, _THEME_FONT_ATTRS)
    faces.set(_w("eastAsia"), face)
    faces.set(_w("cs"), face)
    _drop_attrs(rpr.find(_w("color")), _THEME_ATTRS)


def _heading_colour(size_pt, colors):
    """Large text (LARGE_TEXT_PT or more) takes colors.heading, else primary; smaller needs 4.5:1."""
    if size_pt >= LARGE_TEXT_PT:
        return colors.get("heading") or colors["primary"]
    return small_heading_color(colors)


def _border(ppr, edge, colour, size_and_gap):
    """Replace the paragraph borders on ppr with one single-line edge."""
    from docx.oxml import OxmlElement

    old = ppr.find(_w("pBdr"))
    if old is not None:
        ppr.remove(old)
    borders = OxmlElement("w:pBdr")
    line = OxmlElement(f"w:{edge}")
    for name, value in (("val", "single"), ("sz", size_and_gap[0]),
                        ("space", size_and_gap[1]), ("color", _hex(colour))):
        line.set(_w(name), str(value))
    borders.append(line)
    ppr.insert_element_before(borders, *_PBDR_SUCCESSORS)


def _brand_styles(document, colors, fonts):
    from docx.shared import Pt

    styles = document.styles
    body, heading = fonts["office"], fonts["officeHeading"]
    plan = [("Normal", body, colors["text"], STYLE_SIZES_PT["Normal"], None)]
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        size = STYLE_SIZES_PT[name]
        plan.append((name, heading, _heading_colour(size, colors), size, True))
    plan.append(("Footer", body, colors["text"], FOOTER_PT, None))
    plan += [(name, body, colors["text"], None, None) for name in LIST_STYLES]
    for name, face, colour, size, bold in plan:
        style = styles[name]
        _style_font(style, style.element, face, colour, size, bold)
    _border(styles["Title"].element.get_or_add_pPr(), "left", colors["primary"], TITLE_BAR)
    _border(styles["Heading 1"].element.get_or_add_pPr(), "bottom", colors["divider"], H1_RULE)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)  # clears the rule
    _square_bullets(document, colors["primary"], body)


def _square_bullets(document, colour, face):
    """List Bullet's first level becomes BULLET in colour and face."""
    from docx.oxml import OxmlElement

    ref = document.styles["List Bullet"].element.find(_w("pPr", "numPr", "numId"))
    numbering = document.part.numbering_part.element
    abstract_id = str(numbering.num_having_numId(int(ref.get(_w("val")))).abstractNumId.val)
    abstract = next(a for a in numbering.findall(_w("abstractNum"))
                    if a.get(_w("abstractNumId")) == abstract_id)
    level = next(lvl for lvl in abstract.findall(_w("lvl")) if lvl.get(_w("ilvl")) == "0")
    level.find(_w("lvlText")).set(_w("val"), BULLET)
    old = level.find(_w("rPr"))
    if old is not None:
        level.remove(old)
    rpr = OxmlElement("w:rPr")
    faces, tint = OxmlElement("w:rFonts"), OxmlElement("w:color")
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        faces.set(_w(name), face)
    tint.set(_w("val"), _hex(colour))
    rpr.extend((faces, tint))
    level.append(rpr)


def _sync_legacy_styles(document):
    """Keep the older styles-with-effects part identical to the branded styles part."""
    legacy = _part(document, "/word/stylesWithEffects.xml")
    if legacy is not None:
        legacy._blob = _part(document, "/word/styles.xml").blob


# ---------------------------------------------------------------- body


def _inline(paragraph, text):
    """Add text to a paragraph, turning **spans** into bold runs."""
    position = 0
    for match in _BOLD.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        paragraph.add_run(match.group(1)).bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _item(paragraph, text):
    """List item text: a leading label becomes its own bold run, the rest goes through _inline."""
    match = _LABEL.match(text)
    if match:
        paragraph.add_run(match.group(0)).bold = True
        text = text[match.end():]
    _inline(paragraph, text)


def _new_numbering(document):
    """A fresh numbering instance for the List Number definition, restarting at 1."""
    style_ref = document.styles["List Number"].element.find(_w("pPr", "numPr", "numId"))
    definitions = document.part.numbering_part.element
    source = definitions.num_having_numId(int(style_ref.get(_w("val"))))
    instance = definitions.add_num(source.abstractNumId.val)
    instance.add_lvlOverride(ilvl=0).add_startOverride(1)
    return instance.numId


def _number(paragraph, num_id):
    numpr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numpr.get_or_add_ilvl().val = 0
    numpr.get_or_add_numId().val = num_id


def _render_body(document, markdown_text):
    list_id = None
    for raw in markdown_text.splitlines():
        line = raw.strip()
        numbered = _NUMBERED.fullmatch(line)
        if numbered:
            if list_id is None:
                list_id = _new_numbering(document)
            paragraph = document.add_paragraph(style="List Number")
            _number(paragraph, list_id)
            _item(paragraph, numbered.group(1))
            continue
        list_id = None
        if not line:
            continue
        heading = _HEADING.fullmatch(line)
        bullet = _BULLET.fullmatch(line)
        if heading:
            _inline(document.add_paragraph(style=f"Heading {len(heading.group(1))}"),
                    heading.group(2))
        elif bullet:
            _item(document.add_paragraph(style="List Bullet"), bullet.group(1))
        else:
            paragraph = document.add_paragraph()
            _inline(paragraph, line)
            if _BOLD_LABEL_LINE.fullmatch(line):
                paragraph.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------- title block and footer


def _title_block(document, title, byline, colors, fonts):
    """Title and byline; the byline repeats the Title's left bar so Word draws one bar."""
    document.add_paragraph(title, style="Title")
    paragraph = document.add_paragraph()
    _border(paragraph._p.get_or_add_pPr(), "left", colors["primary"], TITLE_BAR)
    run = paragraph.add_run(byline)
    _style_font(run, run._r, fonts["office"], colors["primary"], BYLINE_PT, bold=True)


def _lockup(brand, brand_path):
    """Path of the brand's logo lockup, or None when the brand has none."""
    name = (brand.get("logo") or {}).get("lockup")
    if not name:
        return None
    path = os.path.join(os.path.dirname(os.path.abspath(brand_path)), name)
    if not os.path.isfile(path):
        raise FileNotFoundError(f"logo lockup named in the brand file not found: {path}")
    return path


def _shade(cell, colour):
    from docx.oxml import OxmlElement

    shading = OxmlElement("w:shd")
    for name, value in (("val", "clear"), ("color", "auto"), ("fill", _hex(colour))):
        shading.set(_w(name), value)
    cell._tc.get_or_add_tcPr().append(shading)


def _first_page_header(document, colors, lockup):
    """Band across the top edge of page 1, then the lockup on the right; later pages bare."""
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Inches, Pt

    section = document.sections[0]
    section.different_first_page_header_footer = True
    section.header_distance = 0
    header = section.first_page_header
    header.is_linked_to_previous = False
    band = header.add_table(1, 2, section.page_width)
    tbl_pr = band._tbl.tblPr
    indent = OxmlElement("w:tblInd")
    indent.set(_w("w"), str(-round(section.left_margin / 635)))
    indent.set(_w("type"), "dxa")
    margins = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        zero = OxmlElement(f"w:{edge}")
        zero.set(_w("w"), "0")
        zero.set(_w("type"), "dxa")
        margins.append(zero)
    fixed = OxmlElement("w:tblLayout")
    fixed.set(_w("type"), "fixed")
    tail = ("w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
    tbl_pr.insert_element_before(indent, "w:tblBorders", "w:shd", "w:tblLayout",
                                 "w:tblCellMar", *tail)
    tbl_pr.insert_element_before(fixed, "w:tblCellMar", *tail)
    tbl_pr.insert_element_before(margins, *tail)
    width = tbl_pr.find(_w("tblW"))
    width.set(_w("type"), "dxa")
    width.set(_w("w"), str(round(section.page_width / 635)))
    row = band.rows[0]
    row.height, row.height_rule = Inches(BAND_IN), WD_ROW_HEIGHT_RULE.EXACTLY
    shares = (BAND_SPLIT, 1 - BAND_SPLIT)
    for column, share in zip(band._tbl.tblGrid.findall(_w("gridCol")), shares):
        column.set(_w("w"), str(round(section.page_width * share / 635)))
    for cell, share, key in zip(row.cells, shares, ("primary", "accent1")):
        cell.width = round(section.page_width * share)
        _shade(cell, colors.get(key) or colors["primary"])
        cell.paragraphs[0].paragraph_format.space_after = 0
    logo = header.paragraphs[0]
    header._element.remove(band._tbl)
    logo._p.addprevious(band._tbl)
    logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    logo.paragraph_format.space_before = Pt(18)
    logo.paragraph_format.space_after = Pt(12)
    logo.add_run().add_picture(lockup, height=Inches(LOCKUP_HEADER_IN))


def _field_runs(paragraph, instruction, styler):
    from docx.oxml import OxmlElement

    for kind in ("begin", "instr", "separate", "result", "end"):
        run = paragraph.add_run("1" if kind == "result" else None)
        styler(run)
        if kind == "instr":
            code = OxmlElement("w:instrText")
            code.set(_XML_SPACE, "preserve")
            code.text = f" {instruction} "
            run._r.append(code)
        elif kind != "result":
            mark = OxmlElement("w:fldChar")
            mark.set(_w("fldCharType"), kind)
            run._r.append(mark)


def _footer_tabs(document, paragraph, section):
    """Clear every inherited tab stop and set one right stop at the text width."""
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Twips

    width = round((section.page_width - section.left_margin - section.right_margin) / 635)
    inherited, style = set(), document.styles["Footer"]
    while style is not None:
        for stop in style.element.findall(_w("pPr", "tabs", "tab")):
            if stop.get(_w("val")) != "clear":
                inherited.add(int(stop.get(_w("pos"))))
        style = style.base_style
    stops = paragraph.paragraph_format.tab_stops
    for position in sorted(inherited - {width}):
        stops.add_tab_stop(Twips(position), WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(Twips(width), WD_TAB_ALIGNMENT.RIGHT)


def _write_footer(document, left_text, colors, fonts, lockup=None):
    """Same footer on every page, including a distinct first page when the section has one."""
    from docx.shared import Inches

    section = document.sections[0]
    footers = [section.footer]
    if section.different_first_page_header_footer:
        footers.append(section.first_page_footer)
    for footer in footers:
        footer.is_linked_to_previous = False
        for old in list(footer._element):
            footer._element.remove(old)
        paragraph = footer.add_paragraph(style="Footer")
        _footer_tabs(document, paragraph, section)

        def plain(run):
            _style_font(run, run._r, fonts["office"], colors["text"], FOOTER_PT)

        def number(run):
            _style_font(run, run._r, fonts["office"], colors["primary"], FOOTER_PT, bold=True)

        for text in (left_text, "\t"):
            plain(paragraph.add_run(text))
        if lockup:
            paragraph.add_run().add_picture(lockup, height=Inches(LOCKUP_FOOTER_IN))
            plain(paragraph.add_run("    "))
        number(paragraph.add_run("Page "))
        _field_runs(paragraph, "PAGE", number)
        number(paragraph.add_run(" of "))
        _field_runs(paragraph, "NUMPAGES", number)


# ---------------------------------------------------------------- build


def _stamp_properties(document, *, title, operator, doc_type, subject, meeting):
    core = document.core_properties
    core.title = f"{title} - {doc_type}"
    core.subject = subject
    core.category = doc_type
    core.keywords = ""
    core.author = operator
    core.comments = f"Prepared by {operator}"
    core.last_modified_by = operator
    core.revision = 1
    core.created = core.modified = meeting


def convert(markdown_text, out, *, title, topic, date, operator, doc_type, type_label,
            footer=None, brand=None):
    """Validate, build the branded document from markdown_text, save it to out; return out."""
    meeting = validate(title=title, topic=topic, date=date, operator=operator,
                       doc_type=doc_type, type_label=type_label)
    brand_path = locate_brand(brand)
    with open(brand_path, encoding="utf-8") as fh:
        brand_data = json.load(fh)
    colors, fonts = brand_data["colors"], brand_data["fonts"]
    lockup = _lockup(brand_data, brand_path)

    from docx import Document

    title, operator = title.strip(), operator.strip()
    doc_type, type_label, topic = doc_type.strip(), type_label.strip(), topic.strip()
    when = format_date(meeting)
    left = footer.strip() if footer and footer.strip() else f"{title}, {when}"

    document = Document()
    body = document.element.body
    for paragraph in body.findall(_w("p")):
        body.remove(paragraph)
    _brand_styles(document, colors, fonts)
    _stamp_properties(document, title=title, operator=operator, doc_type=doc_type,
                      subject=left, meeting=meeting)
    _title_block(document, title, f"{topic} | {when} | {type_label} | Prepared by {operator}",
                 colors, fonts)
    _render_body(document, markdown_text)
    if lockup:
        _first_page_header(document, colors, lockup)
    _write_footer(document, left, colors, fonts, lockup)
    _apply_theme(document, colors, fonts)
    _sync_legacy_styles(document)
    document.save(out)
    return out


def same_path(first, second):
    """True when two paths name the same file, whether or not it exists yet."""
    if os.path.exists(first) and os.path.exists(second):
        return os.path.samefile(first, second)
    return os.path.normcase(os.path.realpath(first)) == os.path.normcase(os.path.realpath(second))


def _say(message):
    print("md2docx: " + " ".join(str(message).split()), file=sys.stderr)


def main(argv=None):
    parser = argparse.ArgumentParser(prog="md2docx.py",
                                     description="Render Markdown notes as a branded .docx.")
    parser.add_argument("notes", help="the Markdown notes file")
    parser.add_argument("--out", required=True, help="the .docx file to write")
    parser.add_argument("--title", required=True, help="the meeting name")
    parser.add_argument("--topic", required=True, help="what the meeting was about, alone")
    parser.add_argument("--date", required=True, help='meeting date, "Month D, YYYY"')
    parser.add_argument("--operator", required=True, help="the person running the session")
    parser.add_argument("--doc-type", required=True, help="document category (properties)")
    parser.add_argument("--type-label", required=True, help="page label (byline)")
    parser.add_argument("--footer", help="left footer text (default: <title>, <date>)")
    parser.add_argument("--brand", help="brand.json to use (see the lookup order)")
    args = parser.parse_args(argv)
    fields = dict(title=args.title, topic=args.topic, date=args.date, operator=args.operator,
                  doc_type=args.doc_type, type_label=args.type_label)
    try:
        validate(**fields)
        if same_path(args.notes, args.out):
            _refuse("--out", "is the notes file itself", "a separate .docx path")
    except ValueError as exc:
        _say(exc)
        return 2
    try:
        with open(args.notes, encoding="utf-8") as fh:
            markdown_text = fh.read()
        convert(markdown_text, args.out, footer=args.footer, brand=args.brand, **fields)
    except ImportError:
        _say("python-docx is not installed; install it with: pip install python-docx")
        return 1
    except KeyError as exc:
        _say(f"build stopped: {exc.args[0] if exc.args else exc}")
        return 1
    except (OSError, ValueError) as exc:
        _say(exc)
        return 1
    print(args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())

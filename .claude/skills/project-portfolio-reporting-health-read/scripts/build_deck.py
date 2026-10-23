"""Build the health read deck from findings.json: a 16:9 pptx in four sections (leadership
summary, team scorecard, fix list per lead, cleanup per person).

    python scripts/build_deck.py <findings.json> --out <deck.pptx> [--brand <path>]

The deck is built from a blank presentation every run. Card capacity is measured from the
layout in deck_layout.py; overflow continues on a new slide (the attention cards roll up into
one line instead), and text never shrinks. After building, every shape is swept against the
slide edges.

Exit 0 written; 1 unreadable or malformed findings, wrong findings_version, an unreadable
brand file, python-pptx missing, or a usage error; 4 layout check failed (nothing written).
Every failure is one stderr line. Needs python-pptx, imported inside functions.
"""
import argparse
import datetime
import io
import os
import re
import sys

sys.dont_write_bytecode = True
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _SCRIPT_DIR)

import deck_layout as lay  # noqa: E402
import health_core as core  # noqa: E402
from deck_layout import LayoutError  # noqa: E402
from health_core import InputError  # noqa: E402

PROG = "build_deck"
HEX = re.compile(r"^#[0-9A-Fa-f]{6}$")
COLOR_KEYS = ("primary", "text", "textMuted", "background", "surface", "divider")
EMU_PER_PT = 12700
HLINK_EXT_URI = "{A12FA001-AC4F-418D-AE19-62706E023703}"
HLINK_NS = "http://schemas.microsoft.com/office/drawing/2018/hyperlinkcolor"
SECTIONS =("Leadership summary", "Team scorecard", "Fix list per lead", "Cleanup per person")
SEVERITY = ("abandoned", "absent", "stale", "conflicting", "misaligned", "repeated",
            "placeholder")
BAND_WORDS = {"current": "current", "aging": "aging", "stale": "stale", "absent": "absent",
              "new": "new", "repeated": "repeated", "placeholder": "placeholder",
              "not_applicable": "not applicable", "aligned": "aligned",
              "misaligned": "misaligned", "not_assessable": "not assessable",
              "not_judged": "not judged", "clean": "clean", "minor": "minor",
              "conflicting": "conflicting"}


class Parser(argparse.ArgumentParser):
    def error(self, message):
        sys.stderr.write("%s: usage error: %s\n" % (PROG, message))
        sys.exit(1)


def parse_args(argv):
    parser = Parser(prog="build_deck.py", description="Build the health read deck.")
    parser.add_argument("findings", help="findings.json written by grade_portfolio.py")
    parser.add_argument("--out", required=True, help="pptx path to write")
    parser.add_argument("--brand", default=None, help="brand JSON path")
    return parser.parse_args(argv)


# ----------------------------------------------------------------- inputs

FINDINGS_KEYS = ("as_of", "title", "prepared_by", "executive_reader_role", "grade_statement",
                 "cleanup_subtitle", "portfolio", "inventory", "splits", "projects", "leads",
                 "unattributed", "corrections_by_project", "projects_needing_nothing",
                 "reporting_well", "closed_since_prior", "method", "context")


def load_findings(path):
    data = core.load_json(path, "findings")
    if not isinstance(data, dict):
        raise InputError("findings %s must hold a JSON object" % path)
    if data.get("findings_version") != 1:
        raise InputError("findings_version must be 1, got %r" % data.get("findings_version"))
    for key in FINDINGS_KEYS:
        if key not in data:
            raise InputError("findings is missing %s" % key)
    return data


def load_brand(path):
    data = core.load_json(path, "brand file")
    if not isinstance(data, dict):
        raise InputError("brand file %s must hold a JSON object" % path)
    colors = data.get("colors") if isinstance(data.get("colors"), dict) else {}
    status = data.get("status") if isinstance(data.get("status"), dict) else {}
    fonts = data.get("fonts") if isinstance(data.get("fonts"), dict) else {}
    brand = {"colors": {}, "status": {}, "fonts": {}}
    for key in COLOR_KEYS:
        value = colors.get(key)
        if not isinstance(value, str) or not HEX.match(value.strip()):
            raise InputError("brand key colors.%s must be #RRGGBB, got %r" % (key, value))
        brand["colors"][key] = value.strip().upper()
    link = colors.get("link")
    if link is None:
        link = brand["colors"]["primary"]
    elif not isinstance(link, str) or not HEX.match(link.strip()):
        raise InputError("brand key colors.link must be #RRGGBB, got %r" % link)
    brand["colors"]["link"] = link.strip().upper()
    for key in ("good", "warning"):
        value = status.get(key)
        if value is not None and (not isinstance(value, str) or not HEX.match(value.strip())):
            raise InputError("brand key status.%s must be #RRGGBB or null, got %r"
                             % (key, value))
        brand["status"][key] = value.strip().upper() if value else None
    office = fonts.get("office")
    if not isinstance(office, str) or not office.strip():
        raise InputError("brand key fonts.office is missing or empty")
    heading = fonts.get("officeHeading")
    brand["fonts"]["office"] = office.strip()
    brand["fonts"]["heading"] = heading.strip() if isinstance(heading, str) and \
        heading.strip() else office.strip()
    return brand


# ----------------------------------------------------------------- color rules

class Palette:
    def __init__(self, brand):
        self.c = brand["colors"]
        self.status = brand["status"]

    def text_on(self, key, bg_key):
        """A text color that reaches 4.5:1 on the background, falling back to colors.text."""
        if core.contrast(self.c[key], self.c[bg_key]) >= 4.5:
            return self.c[key]
        return self.c["text"]

    def link_on(self):
        """One link style for the whole deck: judged on the lowest-contrast background a
        link sits on. PowerPoint always underlines a link, so the file says underline."""
        worst = min(core.contrast(self.c["link"], self.c[key])
                    for key in ("background", "surface"))
        if worst >= 4.5:
            return self.c["link"], True
        return self.c["text"], True

    def chip(self, grade):
        fill = self.status["good"] if grade in ("A", "B") else self.status["warning"]
        fill = fill or self.c["textMuted"]
        label = max((self.c["text"], self.c["background"]),
                    key=lambda color: core.contrast(color, fill))
        return fill, label


# ----------------------------------------------------------------- text blocks

def run(text, link=None, bold=None, color=None):
    return {"text": text, "link": link, "bold": bold, "color": color}


def para(runs, size=lay.SMALL_PT, bold=False, color="text", space_before=0):
    if isinstance(runs, str):
        runs = [run(runs)]
    return {"runs": runs, "size": size, "bold": bold, "color": color,
            "space_before": space_before}


def lower_first(text):
    return text[:1].lower() + text[1:]


def sentence(text):
    text = text.strip()
    return text if text.endswith((".", "!", "?", ")")) else text + "."


def para_text(p):
    return "".join(r["text"] for r in p["runs"])


def paras_height(paras, width):
    return sum(lay.text_height(para_text(p), width, p["size"], p["bold"]) + p["space_before"]
               for p in paras)


def block(name, paras, width=lay.CONTENT_W, fill=None, pad=0, gap=lay.BLOCK_GAP, **extra):
    h = paras_height(paras, width - 2 * pad) + 2 * pad
    item = {"name": name, "paras": paras, "w": width, "fill": fill, "pad": pad, "h": h,
            "gap": gap}
    item.update(extra)
    return item


def text_el(name, x, y, w, h, paras, fill=None, pad=0, anchor="top"):
    return {"type": "text", "name": name, "x": x, "y": y, "w": w, "h": h, "paras": paras,
            "fill": fill, "pad": pad, "anchor": anchor}


def flow_elements(pages_of_blocks, x=lay.MARGIN):
    out = []
    for page in pages_of_blocks:
        els = []
        for b, y in page:
            if b.get("table"):
                els.append(dict(b["table"], y=y))
            else:
                els.append(text_el(b["name"], x, y, b["w"], b["h"], b["paras"], b["fill"],
                                   b["pad"]))
        out.append(els)
    return out


def linked_names(label, items, empty="none"):
    runs = [run(label, bold=True)]
    if not items:
        runs.append(run(empty))
    for i, item in enumerate(items):
        if i:
            runs.append(run("; "))
        runs.append(run(item["name"], link=item.get("url")))
    return runs


def fit_names(prefix, names, width, size, lines, tail):
    """The prefix plus as many names as fit in the given lines; the rest become a count."""
    def text(k):
        shown = "; ".join(names[:k])
        rest = len(names) - k
        if rest:
            return "%s%s%s" % (prefix, shown + " " if shown else "", tail % rest)
        return prefix + shown
    for k in range(len(names), -1, -1):
        if lay.wrap_count(text(k), width, size) <= lines:
            return text(k)
    raise LayoutError("the line '%s' cannot fit its space" % prefix.strip())


# ----------------------------------------------------------------- per-project wording

def top_severity(p):
    c = p["checks"]
    found = []
    if c["currency"]["abandoned"]:
        found.append("abandoned")
    if c["currency"]["band"] in ("absent", "stale"):
        found.append(c["currency"]["band"])
    if c["records"]["band"] == "conflicting":
        found.append("conflicting")
    if c["color"]["band"] == "misaligned":
        found.append("misaligned")
    if c["newness"]["band"] in ("repeated", "placeholder"):
        found.append(c["newness"]["band"])
    if not found:
        return None
    return min(found, key=SEVERITY.index)


def newest_quote(p, findings):
    for request in findings.get("judgment_requests") or []:
        if request["project_id"] == p["project_id"] and request.get("entry_text"):
            return request["entry_text"]
    return None


def condition(p, severity):
    c = p["checks"]
    if severity == "abandoned":
        return "Set up and never used"
    if severity == "absent":
        return "No status entries"
    if severity == "stale":
        return "Newest status entry is %d days old" % c["currency"]["age_days"]
    if severity == "conflicting":
        return "Records disagree about where the project sits"
    if severity == "misaligned":
        if c["color"]["basis"] == "rule":
            return "A recorded hold shows %s" % c["color"]["reported"]
        return "The reported color is not supported by the written status"
    if severity == "repeated":
        return "The newest entry repeats earlier ones"
    return "The newest entry is a placeholder"


def evidence(p, severity, fixes):
    c = p["checks"]
    if severity == "abandoned":
        return ("No task has left the first status, the status task has not advanced and no "
                "task carries a date.")
    if severity == "absent":
        return "The status task has no entries from a person."
    if severity == "stale":
        return "Last entry %s by %s." % (c["currency"]["latest_entry_date"],
                                         c["currency"]["author"])
    if severity == "conflicting":
        for item in c["records"]["items"]:
            if item["counted"] and item["severity"] == "conflict":
                return item["detail"]
    if severity == "misaligned":
        if c["color"]["quote"]:
            return "The name shows %s; the entry says: '%s'" % (c["color"]["reported"],
                                                               c["color"]["quote"])
        return "The name records a hold and shows %s." % c["color"]["reported"]
    if severity == "repeated":
        if c["newness"]["basis"] == "judgment":
            return "The newest entry repeats the open item from the previous one."
        return "%s last %d entries match once dates are set aside." % (
            "At least the" if c["newness"]["at_least"] else "The",
            c["newness"]["consecutive_same"])
    for f in fixes:
        if f["project_id"] == p["project_id"] and f["kind"] == "placeholder":
            return f["what"].split(". ", 1)[1]
    return "The newest entry is a placeholder."


def short_evidence(p, fixes):
    sev = top_severity(p)
    if sev:
        return evidence(p, sev, fixes)
    c = p["checks"]
    if c["currency"]["band"] == "aging":
        return "Newest entry %d days old." % c["currency"]["age_days"]
    counted = [i for i in c["records"]["items"] if i["counted"]]
    if counted:
        return counted[0]["detail"]
    if c["color"]["band"] == "not_judged":
        return "Color judgment still missing."
    return "Current, new and consistent; records agree."


def latest_cell(p):
    cur = p["checks"]["currency"]
    if not cur["latest_entry_date"]:
        return "None"
    days = cur["age_days"]
    return "%s (%d day%s)" % (cur["latest_entry_date"], days, "" if days == 1 else "s")


def newness_cell(p):
    n = p["checks"]["newness"]
    if n["band"] == "repeated" and n["consecutive_same"] > 1:
        return "Repeated (%d%s)" % (n["consecutive_same"], "+" if n["at_least"] else "")
    return BAND_WORDS[n["band"]].capitalize()


# ----------------------------------------------------------------- slide composition

def slide(section, title, elements, notes, continued=False):
    return {"section": section, "title": title + (" (continued)" if continued else ""),
            "elements": elements, "notes": notes}


def all_fixes(findings):
    out = list(findings["unattributed"])
    for lead in findings["leads"]:
        out += lead["fix_list"]
    return out


def grade_slide(f, pal):
    portfolio = f["portfolio"]
    as_of = "As of %s." % f["as_of"]
    grade = portfolio["grade"] or "-"
    fill, label = pal.chip(grade)
    chip_paras = [para([run(grade)], size=lay.GRADE_PT, bold=True, color=label)]
    if portfolio["provisional"]:
        chip_paras.append(para([run("(provisional)")], size=lay.BODY_PT, color=label))
    els = [{"type": "chip", "name": "grade-chip", "x": lay.MARGIN, "y": lay.BODY_TOP,
            "w": lay.CHIP_W, "h": lay.CHIP_H, "paras": chip_paras, "fill_hex": fill}]
    splits = f["splits"]
    lines = []
    for label_text, key in (("Currency", "currency"), ("Newness", "newness"),
                            ("Color", "color"), ("Records", "records")):
        parts = ", ".join("%d %s" % (v, BAND_WORDS[k]) for k, v in splits[key].items())
        lines.append(para("%s: %s" % (label_text, parts), size=lay.BODY_PT,
                          space_before=4))
    headline = portfolio["reasons"][0] if portfolio["reasons"] else "No projects in scope."
    y = lay.BODY_TOP
    head_paras = [para(headline[0].upper() + headline[1:] + ".", size=lay.HEADING_PT,
                       bold=True)]
    h = paras_height(head_paras, lay.RIGHT_W)
    els.append(text_el("body-headline", lay.RIGHT_X, y, lay.RIGHT_W, h, head_paras))
    y += h + lay.BLOCK_GAP
    h = paras_height(lines, lay.RIGHT_W)
    els.append(text_el("body-checks", lay.RIGHT_X, y, lay.RIGHT_W, h, lines))
    y += h
    if portfolio.get("reasons_changed"):
        changed = [para("The letter is unchanged from the prior read; the reasons changed.",
                        size=lay.SMALL_PT, space_before=4)]
        h = paras_height(changed, lay.RIGHT_W)
        y += lay.BLOCK_GAP
        els.append(text_el("body-unchanged", lay.RIGHT_X, y, lay.RIGHT_W, h, changed))
        y += h
    right_bottom = max(y, lay.BODY_TOP + lay.CHIP_H)

    cover_text = " | ".join([
        "As of %s" % f["as_of"],
        "Prior read %s" % f["prior_as_of"] if f.get("prior_as_of") else "No prior read supplied",
        "Prepared for the %s" % f["executive_reader_role"],
        "Prepared by %s" % f["prepared_by"]])
    cover = [para(cover_text, size=lay.SMALL_PT, color="textMuted")]
    cover_h = paras_height(cover, lay.CONTENT_W)
    cover_y = lay.BODY_BOTTOM - cover_h
    half = (lay.CONTENT_W - 3 * lay.PANEL_PAD) / 2
    gs = f["grade_statement"]
    parts = []
    for name, label_text, text in (("body-measures", "What this grade measures",
                                    gs["measures"]),
                                   ("body-not-measured", "What it does not measure",
                                    gs["does_not_measure"])):
        ps = [para([run(label_text)], size=lay.SMALL_PT, bold=True),
              para(text, size=lay.BODY_PT, space_before=2)]
        parts.append((name, ps, paras_height(ps, half)))
    panel_h = max(p[2] for p in parts) + 2 * lay.PANEL_PAD
    panel_y = cover_y - lay.PANEL_GAP - panel_h
    if panel_y < right_bottom + lay.PANEL_GAP:
        raise LayoutError("the grade slide text does not fit above its panel")
    els.append({"type": "rect", "name": "panel", "x": lay.MARGIN, "y": panel_y,
                "w": lay.CONTENT_W, "h": panel_h, "fill": "surface"})
    for i, (name, ps, h) in enumerate(parts):
        x = lay.MARGIN + lay.PANEL_PAD + i * (half + lay.PANEL_PAD)
        els.append(dict(text_el(name, x, panel_y + lay.PANEL_PAD, half, h, ps), bg="surface"))
    els.append(text_el("caption-cover", lay.MARGIN, cover_y, lay.CONTENT_W, cover_h, cover))
    notes = [as_of, "Portfolio grade %s%s." % (grade, " (provisional)"
                                                if portfolio["provisional"] else "")]
    notes += ["Reason: %s." % r for r in portfolio["reasons"]]
    notes.append(portfolio["points_note"])
    return [slide(SECTIONS[0], "Reporting reliability grade", els, notes)]


def attention_slide(f, pal, fixes):
    as_of = "As of %s." % f["as_of"]
    exceptions = []
    for p in f["projects"]:
        sev = top_severity(p)
        if sev:
            exceptions.append((SEVERITY.index(sev), len(exceptions), p, sev))
    exceptions.sort(key=lambda e: (e[0], e[1]))
    cap = lay.attention_capacity()
    shown, rest = exceptions[:cap], exceptions[cap:]
    els = []
    inner = lay.CONTENT_W - 2 * lay.CARD_PAD
    name_cpl = lay.chars_per_line(inner, lay.CARD_NAME_PT, bold=True)
    line_cpl = lay.chars_per_line(inner, lay.CARD_LINE_PT)
    for i, (_, _, p, sev) in enumerate(shown):
        if len(p["name"]) > name_cpl:
            raise LayoutError("the project name %r is longer than the %d characters a card "
                              "holds on one line; shorten the list name in the tracker and "
                              "pull again" % (p["name"][:40] + "...", name_cpl))
        ev = evidence(p, sev, fixes)
        if len(ev) > line_cpl:
            ev = ev[:line_cpl - 4].rsplit(" ", 1)[0] + " ..."
        ps = [para([run(p["name"], link=p["url"])], size=lay.CARD_NAME_PT, bold=True),
              para(condition(p, sev), size=lay.CARD_LINE_PT),
              para(ev, size=lay.CARD_LINE_PT, color="textMuted")]
        y = lay.CARD_TOP + i * (lay.CARD_H + lay.CARD_GAP)
        els.append(dict(text_el("card-%d" % (i + 1), lay.MARGIN, y, lay.CONTENT_W, lay.CARD_H,
                                ps, fill="surface", pad=lay.CARD_PAD), bg="surface"))
    if not exceptions:
        ps = [para("No in-scope project needs attention this read.", size=lay.BODY_PT)]
        els.append(text_el("body-none", lay.MARGIN, lay.CARD_TOP, lay.CONTENT_W,
                           paras_height(ps, lay.CONTENT_W), ps))
    y = lay.ATTENTION_BAND_BOTTOM + lay.PANEL_GAP
    if rest:
        names = [e[2]["name"] for e in rest]
        text = fit_names("Also needing attention (%d): " % len(rest), names, lay.CONTENT_W,
                         lay.SMALL_PT, 2, "and %d more; see Project scorecard")
        els.append(text_el("rollup", lay.MARGIN, y, lay.CONTENT_W, lay.FOOT_LINE_H,
                           [para(text, size=lay.SMALL_PT)]))
    y += lay.FOOT_LINE_H + lay.PANEL_GAP
    well = f["reporting_well"]
    text = fit_names("Reporting well: ", well, lay.CONTENT_W, lay.SMALL_PT, 2,
                     "and %d more") if well else "Reporting well: no project this read."
    els.append(text_el("body-closing", lay.MARGIN, y, lay.CONTENT_W, lay.FOOT_LINE_H,
                       [para(text, size=lay.SMALL_PT, color="textMuted")]))
    notes = [as_of, "Exceptions in severity order (%d):" % len(exceptions)]
    notes += ["%s: %s." % (e[2]["name"], condition(e[2], e[3])) for e in exceptions]
    return [slide(SECTIONS[0], "Needs attention", els, notes)]


def stats_block(name, title, values):
    head = block(name + "-head", [para([run(title)], size=lay.SMALL_PT, bold=True)],
                 keep_next=True, gap=10)
    box_h = 20 * lay.LINE_FACTOR + lay.CAPTION_PT * lay.LINE_FACTOR + 8
    width = (lay.CONTENT_W - 3 * 10) / 4.0
    cells = []
    for i, (label, value) in enumerate(values):
        cells.append(text_el("stat-%s-%d" % (name, i + 1), lay.MARGIN + i * (width + 10), 0,
                             width, box_h,
                             [para([run(str(value))], size=20, bold=True),
                              para([run(label)], size=lay.CAPTION_PT, color="textMuted")],
                             fill="surface", pad=4))
    stats = {"name": name, "h": box_h, "gap": 4, "table": None, "cells": cells}
    return [head, stats]


def snapshot_slides(f, pal):
    inv = f["inventory"]
    blocks = [
        block("body-active", [para(linked_names("Active (%d): " % len(inv["active"]),
                                                inv["active"]))]),
        block("body-completed", [para(linked_names("Completed (%d): " % len(inv["completed"]),
                                                   inv["completed"]))]),
        block("body-not-started",
              [para(linked_names("Not started (%d): " % len(inv["not_started"]),
                                 inv["not_started"]))]),
    ]
    sp = f["splits"]
    blocks += stats_block("currency", "Status currency (in-scope projects)",
                          [(BAND_WORDS[k].capitalize(), v) for k, v in sp["currency"].items()])
    blocks += stats_block("color", "Color alignment (in-scope projects)",
                          [(BAND_WORDS[k].capitalize(), v) for k, v in sp["color"].items()])
    held = "; ".join(h["name"] for h in inv["held"]) or "none"
    blocks.append(block("body-held", [para([run("Out of scope: on hold (%d): " % len(inv["held"]),
                                                bold=True), run(held)])], gap=10))
    if inv["hold_pending"]:
        names = "; ".join(h["name"] for h in inv["hold_pending"])
        blocks.append(block("body-pending",
                            [para([run("Awaiting hold judgment (%d): " % len(inv["hold_pending"]),
                                       bold=True), run(names)])]))
    pages = lay.place_blocks(blocks, lay.BODY_TOP, lay.BODY_BOTTOM)
    out = []
    for i, page in enumerate(pages):
        els = []
        for b, y in page:
            if "cells" in b:
                for cell in b["cells"]:
                    els.append(dict(cell, y=y, bg="surface"))
            else:
                els.append(text_el(b["name"], lay.MARGIN, y, b["w"], b["h"], b["paras"]))
        notes = ["As of %s." % f["as_of"],
                 "Active %d, completed %d, not started %d; on hold %d, awaiting hold judgment %d."
                 % (len(inv["active"]), len(inv["completed"]), len(inv["not_started"]),
                    len(inv["held"]), len(inv["hold_pending"]))]
        out.append(slide(SECTIONS[0], "Portfolio snapshot", els, notes, continued=i > 0))
    return out


def table_slides(section, title, columns, rows, f, notes_lines, captions=(), closing=None,
                 empty_text="Nothing to show."):
    """rows: list of lists of cells; a cell is a list of paras. Paginates on estimated heights
    with the header repeated, then places the closing blocks after the table."""
    widths = [w for _, w in columns]
    header_cells = [[para([run(name)], size=lay.TABLE_PT, bold=True)] for name, _ in columns]

    def cell_text(cell):
        return "\n".join(para_text(p) for p in cell)

    def height(cells):
        return lay.row_height([cell_text(c) for c in cells], widths, lay.TABLE_PT)

    header_h = lay.row_height([name for name, _ in columns], widths, lay.TABLE_PT) + 2
    bottom = lay.TABLE_BAND_BOTTOM if captions else lay.BODY_BOTTOM
    heights = [height(r) for r in rows]
    pages = lay.paginate_rows(heights, header_h, lay.BODY_TOP, bottom) if rows else [[]]
    out = []
    as_of = "As of %s." % f["as_of"]
    for i, idx in enumerate(pages):
        els = []
        if rows:
            table_rows = [{"h": header_h, "header": True, "cells": header_cells}]
            table_rows += [{"h": heights[j], "header": False, "cells": rows[j]} for j in idx]
            els.append({"type": "table", "name": "table", "x": lay.MARGIN, "y": lay.BODY_TOP,
                        "widths": widths, "rows": table_rows})
        else:
            ps = [para(empty_text, size=lay.BODY_PT)]
            els.append(text_el("body-empty", lay.MARGIN, lay.BODY_TOP, lay.CONTENT_W,
                               paras_height(ps, lay.CONTENT_W), ps))
        cap_y = lay.BODY_BOTTOM - lay.CAPTION_H
        if captions:
            ps = [para(c, size=lay.CAPTION_PT, color="textMuted") for c in captions]
            els.append(text_el("caption-table", lay.MARGIN, cap_y, lay.CONTENT_W,
                               lay.CAPTION_H, ps))
        out.append([els, sum(heights[j] for j in idx) + header_h if rows else 40])
    slides = [slide(section, title, page[0], [as_of] + notes_lines, continued=i > 0)
              for i, page in enumerate(out)]
    if closing:
        used = lay.BODY_TOP + out[-1][1] + lay.PANEL_GAP
        pages = lay.place_blocks(closing, used, bottom)
        slides[-1]["elements"] += flow_elements(pages[:1])[0]
        for extra in flow_elements(pages[1:]):
            slides.append(slide(section, title, extra, [as_of] + notes_lines, continued=True))
    return slides


def scorecard_slides(f, pal, fixes):
    rows, notes = [], []
    for p in f["projects"]:
        c = p["checks"]
        grade = p["grade"] + ("*" if p["provisional"] else "")
        cells = [p["name"], p["lead"] or "Unattributed", latest_cell(p),
                 c["color"]["reported"] or "-", BAND_WORDS[c["color"]["band"]].capitalize(),
                 newness_cell(p), BAND_WORDS[c["records"]["band"]].capitalize(), grade,
                 short_evidence(p, fixes)]
        row = [[para([run(cells[0], link=p["url"])], size=lay.TABLE_PT)]]
        row += [[para(text, size=lay.TABLE_PT)] for text in cells[1:]]
        rows.append(row)
        notes.append("%s: grade %s, %d points." % (p["name"], grade, p["points"]))
    captions = []
    if any(p["provisional"] for p in f["projects"]):
        captions.append("* Provisional: a judgment is still missing for this project.")
    captions.append("Context: %d overdue tasks across in-scope lists. Not graded."
                    % f["context"]["overdue_tasks_total"])
    return table_slides(SECTIONS[1], "Project scorecard", lay.SCORECARD_COLUMNS, rows, f,
                        notes, captions=captions, empty_text="No projects are in scope.")


def detail_slides(f, pal, fixes):
    groups = (("Not current", []), ("Repeated or placeholder", []), ("Color conflicts", []))
    for p in f["projects"]:
        c = p["checks"]
        if c["currency"]["result"] in ("minor", "fail"):
            if c["currency"]["abandoned"]:
                why = evidence(p, "abandoned", fixes)
                text = "set up and never used; %s" % (why[0].lower() + why[1:])
            elif c["currency"]["band"] == "absent":
                text = "no status entries from a person"
            else:
                text = "newest entry %s, %d days old (%s)" % (
                    c["currency"]["latest_entry_date"], c["currency"]["age_days"],
                    c["currency"]["band"])
            groups[0][1].append((p, text))
        if c["newness"]["band"] in ("repeated", "placeholder"):
            groups[1][1].append((p, evidence(p, c["newness"]["band"], fixes)))
        if c["color"]["band"] == "misaligned":
            groups[2][1].append((p, evidence(p, "misaligned", fixes)))
    blocks = []
    for gi, (title, items) in enumerate(groups):
        blocks.append(block("body-group-head", [para([run(title)], size=lay.HEADING_PT,
                                                     bold=True)],
                            keep_next=True, gap=10 if gi else 0))
        if not items:
            blocks.append(block("body-group-none", [para("None this read.")]))
        for p, text in items:
            blocks.append(dict(block("card-detail", [para([run(p["name"], bold=True),
                                                           run(": " + text)])],
                                     fill="surface", pad=6), bg="surface"))
    well = f["reporting_well"]
    blocks.append(block("body-closing", [para("Reporting well: %s" % ("; ".join(well) if well
                                                                     else "no project this read."),
                                              color="textMuted")], gap=10))
    pages = lay.place_blocks(blocks, lay.BODY_TOP, lay.BODY_BOTTOM)
    out = []
    for i, els in enumerate(flow_elements(pages)):
        for el in els:
            if el["fill"] == "surface":
                el["bg"] = "surface"
        notes = ["As of %s." % f["as_of"]] + ["%s: %s" % (t, len(items)) for t, items in groups]
        out.append(slide(SECTIONS[1], "Status and color detail", els, notes, continued=i > 0))
    return out


def method_slides(f, pal):
    m = f["method"]
    filters = m["filters"]
    sections = [
        ("Sources", m["sources"] + ["Pulled at %s; subtasks %s." % (
            m["pulled_at"], "included" if m["includes_subtasks"] else "left out")]),
        ("Filters", ["Tasks past the parked limit: %d raw, %d counted as parked."
                     % (filters["parked_raw"], filters["parked_counted"]),
                     "Status tasks left out of every bucket: %d."
                     % filters["status_tasks_excluded"],
                     "Finished tasks left out of every bucket: %d."
                     % filters["finished_tasks_excluded"]]),
        ("Pairs and roadmap", ["Lists with no pair: %s." % ("; ".join(m["unpaired_active_lists"])
                                                            or "none"),
                               "Unconfirmed pairs (shown, not counted): %s."
                               % ("; ".join(m["unconfirmed_pairs"]) or "none")]
         + ["Roadmap row not found: %s." % x for x in m["missing_roadmap_rows"]]
         + ["Finished project, open roadmap row: %s." % x for x in m["completed_pair_conflicts"]]),
        ("Judgments", ["Not judged: %s." % x for x in m["not_judged"]]
         + ["Set aside: %s." % x for x in m["discarded_judgments"]]),
        ("Counts and changes", m["count_checks"] + m["measurement_changes"]),
        ("Notes", m["notes"]),
        ("What the grade measures", [f["grade_statement"]["measures"],
                                     "It does not measure: " +
                                     lower_first(f["grade_statement"]["does_not_measure"]),
                                     f["portfolio"]["points_note"]]),
    ]
    blocks = []
    for gi, (title, lines) in enumerate(sections):
        group = "method-%d" % gi
        blocks.append(block("body-method-head", [para([run(title)], size=lay.BODY_PT,
                                                      bold=True)],
                            keep_next=True, group=group, gap=10 if gi else 0))
        for line in lines or ["None."]:
            blocks.append(block("body-method-line", [para(sentence(line))], group=group,
                                gap=2))
    pages = lay.place_blocks(blocks, lay.BODY_TOP, lay.BODY_BOTTOM)
    return [slide(SECTIONS[1], "Method, sources and limits", els,
                  ["As of %s." % f["as_of"], "Sources, filters, pairs, judgments and limits."],
                  continued=i > 0) for i, els in enumerate(flow_elements(pages))]


def fix_line(item, prefix=""):
    return para([run(prefix), run(item["project_name"], link=item.get("url")),
                 run(": " + item["what"])])


def fix_list_slides(f, pal):
    owners = [(lead["name"], lead["fix_list"]) for lead in f["leads"]]
    if f["unattributed"]:
        owners.append(("Unattributed", f["unattributed"]))
    blocks = []
    for gi, (owner, items) in enumerate(owners):
        blocks.append(block("body-owner-head", [para([run("%s (%d)" % (owner, len(items)))],
                                                     size=lay.HEADING_PT, bold=True)],
                            keep_next=True, gap=10 if gi else 0, head=True))
        if not items:
            blocks.append(block("body-fix-none", [para("Nothing to fix this read.")]))
        for item in items:
            blocks.append(block("body-fix", [fix_line(item)], gap=3))
    blocks.append(block("body-closed-head", [para([run("Closed since the prior read")],
                                                  size=lay.HEADING_PT, bold=True)],
                        keep_next=True, gap=12))
    if not f.get("prior_as_of"):
        blocks.append(block("body-closed", [para("No prior read supplied.")]))
    elif not f["closed_since_prior"]:
        blocks.append(block("body-closed", [para("None.")]))
    for item in f["closed_since_prior"]:
        blocks.append(block("body-closed", [para("%s, %s: %s" % (
            item.get("owner") or "Unattributed", item.get("project_name"), item.get("what")))],
            gap=3))
    pages = lay.place_blocks(blocks, lay.BODY_TOP, lay.BODY_BOTTOM)
    notes = ["As of %s." % f["as_of"]] + ["%s: %d items." % (o, len(i)) for o, i in owners]
    return [slide(SECTIONS[2], "Fix list by lead", els, notes, continued=i > 0)
            for i, els in enumerate(flow_elements(pages))]


def corrections_slides(f, pal):
    rows = []
    for entry in f["corrections_by_project"]:
        many = len(entry["owners"]) > 1
        corr = [para(("%s: " % c["owner"] if many else "") + c["what"], size=lay.TABLE_PT,
                     space_before=2 if k else 0)
                for k, c in enumerate(entry["corrections"])]
        rows.append([[para([run(entry["name"], link=entry["url"])], size=lay.TABLE_PT)],
                     [para(", ".join(entry["owners"]), size=lay.TABLE_PT)], corr])
    names = f["projects_needing_nothing"]
    closing = [block("body-nothing", [para("Nothing needed this read: %s"
                                           % ("; ".join(names) if names else "none"),
                                           color="textMuted")], gap=0)]
    notes = ["%d projects or lists need a correction." % len(rows)]
    return table_slides(SECTIONS[2], "Corrections by project", lay.CORRECTION_COLUMNS, rows,
                        f, notes, closing=closing, empty_text="No corrections this read.")


def cleanup_slides(f, pal):
    out = []
    subtitle = [para(f["cleanup_subtitle"], size=lay.SMALL_PT, color="textMuted")]
    sub_h = paras_height(subtitle, lay.CONTENT_W)
    top = lay.BODY_TOP + sub_h + 10
    for lead in f["leads"]:
        cl = lead["cleanup"]
        blocks = [block("body-whattofix-head",
                        [para([run("What to fix (%d)" % cl["what_to_fix_count"])],
                              size=lay.HEADING_PT, bold=True)], keep_next=True, gap=0)]
        if cl["clean"]:
            blocks.append(block("body-clean", [para(
                "There is nothing to correct this read. Your lists match the record.",
                size=lay.BODY_PT)]))
        for i, line in enumerate(cl["summary"], 1):
            blocks.append(block("body-summary", [para([
                run("%d. " % i), run(line["project_name"], link=line["url"]),
                run(": " + line["what"])])], gap=3))
        heads = {}
        for bucket in cl["buckets"]:
            key = bucket["key"]
            count = len(bucket["lines"])
            label = "%s (%d)" % (bucket["label"], count)
            explain = "%s %s" % (bucket["definition"], bucket["remedy"])
            heads[key] = (label, explain)
            blocks.append(block("body-bucket-%s" % key, [para([run(label)], size=lay.BODY_PT,
                                                              bold=True)],
                                keep_next=True, gap=12, context=key, head=True))
            blocks.append(block("body-bucket-text", [para(explain, color="textMuted")],
                                keep_next=True, gap=2, context=key, head=True))
            for line in bucket["lines"]:
                name = line["task_name"]
                if line.get("task_id"):
                    name = "%s (%s)" % (name, line["task_id"])
                blocks.append(block("body-line", [para([
                    run(name + ", "), run(line["project_name"], link=line["url"]),
                    run(": " + line["detail"])])], gap=2, context=key))

        def repeat(key, heads=heads):
            label, explain = heads[key]
            return [block("body-bucket-%s" % key, [para([run(label + " (continued)")],
                                                        size=lay.BODY_PT, bold=True)], gap=0),
                    block("body-bucket-text", [para(explain, color="textMuted")], gap=2)]

        pages = lay.place_blocks(blocks, top, lay.BODY_BOTTOM, repeat=repeat)
        for i, els in enumerate(flow_elements(pages)):
            els.insert(0, text_el("body-subtitle", lay.MARGIN, lay.BODY_TOP, lay.CONTENT_W,
                                  sub_h, subtitle))
            notes = ["As of %s." % f["as_of"],
                     "%s: %d things to fix, %d bucket lines." % (
                         lead["name"], cl["what_to_fix_count"],
                         sum(len(b["lines"]) for b in cl["buckets"]))]
            out.append(slide(SECTIONS[3], "Cleanup: %s" % lead["name"], els, notes,
                             continued=i > 0))
    return out


def name_strings(node):
    """Every name taken from the tracker (project, list or task names). They are text the
    build did not write, so the cross-reference guard skips them."""
    found = set()
    if isinstance(node, dict):
        for key, value in node.items():
            if key.endswith("name") and isinstance(value, str) and value:
                found.add(value)
            else:
                found |= name_strings(value)
    elif isinstance(node, list):
        for value in node:
            found |= name_strings(value)
    return found


def compose(f, pal):
    fixes = all_fixes(f)
    slides = []
    slides += grade_slide(f, pal)
    slides += attention_slide(f, pal, fixes)
    slides += snapshot_slides(f, pal)
    slides += scorecard_slides(f, pal, fixes)
    slides += detail_slides(f, pal, fixes)
    slides += method_slides(f, pal)
    slides += fix_list_slides(f, pal)
    slides += corrections_slides(f, pal)
    slides += cleanup_slides(f, pal)
    counts = {s: sum(1 for x in slides if x["section"] == s) for s in SECTIONS}
    seen = {s: 0 for s in SECTIONS}
    for s in slides:
        seen[s["section"]] += 1
        s["footer"] = "%s | page %d of %d" % (s["section"], seen[s["section"]],
                                              counts[s["section"]])
    names = sorted(name_strings(f), key=len, reverse=True)
    for s in slides:
        text = " ".join([s["title"]] + s["notes"])
        for name in names:
            text = text.replace(name, " ")
        if re.search(r"(?i)\bslide \d+", text):
            raise LayoutError("slide text refers to a slide by number")
    return slides


# ----------------------------------------------------------------- pptx

def import_pptx():
    try:
        import pptx  # noqa: F401
    except ImportError:
        raise InputError("python-pptx is required; install it with: pip install python-pptx")


def render(slides, f, brand, pal):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt
    from lxml import etree

    colors = brand["colors"]
    office, heading = brand["fonts"]["office"], brand["fonts"]["heading"]
    prs = Presentation()
    prs.slide_width = Emu(lay.SLIDE_W * EMU_PER_PT)
    prs.slide_height = Emu(lay.SLIDE_H * EMU_PER_PT)
    layouts = list(prs.slide_layouts)
    blank = [l for l in layouts if l.name.lower() == "blank"]
    layout = blank[0] if blank else min(layouts, key=lambda l: len(l.placeholders))

    def emu(value):
        return Emu(int(round(value * EMU_PER_PT)))

    def rgb(hex_color):
        return RGBColor.from_string(hex_color[1:])

    def strip_style(shape):
        style = shape._element.find(qn("p:style"))
        if style is not None:
            shape._element.remove(style)

    def write_paras(frame, paras, bg_key, font):
        for i, p in enumerate(paras):
            target = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            if p["space_before"]:
                target.space_before = Pt(p["space_before"])
            for r in p["runs"]:
                piece = target.add_run()
                piece.text = r["text"]
                piece.font.size = Pt(p["size"])
                piece.font.name = font
                piece.font.bold = bool(r["bold"] if r["bold"] is not None else p["bold"])
                if r["link"]:
                    color, underline = pal.link_on()
                    piece.font.color.rgb = rgb(color)
                    piece.font.underline = underline
                    piece.hyperlink.address = r["link"]
                    own_color_link(piece)
                else:
                    key = r["color"] or p["color"]
                    if key.startswith("#"):
                        piece.font.color.rgb = rgb(key)
                    else:
                        piece.font.color.rgb = rgb(pal.text_on(key, bg_key))

    def own_color_link(piece):
        """Tell PowerPoint to draw the link in the run's own color, not the theme's."""
        click = piece._r.get_or_add_rPr().find(qn("a:hlinkClick"))
        if click is None:
            return
        ext_list = click.makeelement(qn("a:extLst"), {})
        ext = ext_list.makeelement(qn("a:ext"), {"uri": HLINK_EXT_URI})
        color = etree.SubElement(ext, "{%s}hlinkClr" % HLINK_NS, nsmap={"ahyp": HLINK_NS})
        color.set("val", "tx")
        ext_list.append(ext)
        click.append(ext_list)

    def fill_text(frame, paras, bg_key, font, pad=0, anchor="top"):
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = emu(pad)
        frame.margin_top = frame.margin_bottom = emu(pad)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE if anchor == "middle" else MSO_ANCHOR.TOP
        write_paras(frame, paras, bg_key, font)

    def textbox(sl, el, font=office):
        shape = sl.shapes.add_textbox(emu(el["x"]), emu(el["y"]), emu(el["w"]), emu(el["h"]))
        shape.name = el["name"]
        bg = el.get("bg") or ("surface" if el.get("fill") == "surface" else "background")
        if el.get("fill"):
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(colors[el["fill"]])
        fill_text(shape.text_frame, el["paras"], bg, font, el.get("pad", 0),
                  el.get("anchor", "top"))
        return shape

    for spec in slides:
        sl = prs.slides.add_slide(layout)
        for placeholder in list(sl.placeholders):
            placeholder._element.getparent().remove(placeholder._element)
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = rgb(colors["background"])
        label_color = pal.text_on("primary", "background")
        textbox(sl, text_el("section-label", lay.MARGIN, lay.LABEL_TOP, lay.CONTENT_W,
                            lay.LABEL_H, [para([run(spec["section"])], size=lay.LABEL_PT,
                                               bold=True, color=label_color)]))
        textbox(sl, text_el("title", lay.MARGIN, lay.TITLE_TOP, lay.CONTENT_W, lay.TITLE_H,
                            [para([run(spec["title"])], size=lay.TITLE_PT)]), font=heading)
        divider = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(lay.MARGIN),
                                      emu(lay.DIVIDER_TOP), emu(lay.CONTENT_W), emu(1))
        divider.name = "divider"
        strip_style(divider)
        divider.fill.solid()
        divider.fill.fore_color.rgb = rgb(colors["divider"])
        divider.line.fill.background()
        for el in spec["elements"]:
            if el["type"] == "text":
                textbox(sl, el)
            elif el["type"] == "rect":
                shape = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(el["x"]), emu(el["y"]),
                                            emu(el["w"]), emu(el["h"]))
                shape.name = el["name"]
                strip_style(shape)
                shape.fill.solid()
                shape.fill.fore_color.rgb = rgb(colors[el["fill"]])
                shape.line.fill.background()
            elif el["type"] == "chip":
                shape = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(el["x"]), emu(el["y"]),
                                            emu(el["w"]), emu(el["h"]))
                shape.name = el["name"]
                strip_style(shape)
                shape.fill.solid()
                shape.fill.fore_color.rgb = rgb(el["fill_hex"])
                shape.line.fill.background()
                fill_text(shape.text_frame, el["paras"], "background", office, 6, "middle")
                for p in shape.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
            elif el["type"] == "table":
                rows, widths = el["rows"], el["widths"]
                total_h = sum(r["h"] for r in rows)
                frame = sl.shapes.add_table(len(rows), len(widths), emu(el["x"]), emu(el["y"]),
                                            emu(sum(widths)), emu(total_h))
                frame.name = el["name"]
                table = frame.table
                table.first_row = False
                table.horz_banding = False
                for j, w in enumerate(widths):
                    table.columns[j].width = emu(w)
                for i, r in enumerate(rows):
                    table.rows[i].height = emu(r["h"])
                    bg = "surface" if r["header"] else "background"
                    for j, cell_paras in enumerate(r["cells"]):
                        cell = table.cell(i, j)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = rgb(colors[bg])
                        cell.margin_left = cell.margin_right = emu(lay.CELL_PAD_H)
                        cell.margin_top = cell.margin_bottom = emu(lay.CELL_PAD_V)
                        cell.vertical_anchor = MSO_ANCHOR.TOP
                        cell.text_frame.word_wrap = True
                        write_paras(cell.text_frame, cell_paras, bg, office)
        textbox(sl, text_el("footer", lay.MARGIN, lay.FOOTER_TOP, lay.CONTENT_W, lay.FOOTER_H,
                            [para([run(spec["footer"])], size=lay.CAPTION_PT,
                                  color="textMuted")]))
        sl.notes_slide.notes_text_frame.text = "\n".join(spec["notes"])
    props = prs.core_properties
    props.title = f["title"]
    props.author = f["prepared_by"]
    props.last_modified_by = f["prepared_by"]
    props.subject = ""
    props.keywords = ""
    props.comments = ""
    props.revision = 1
    stamp = datetime.datetime.combine(datetime.date.fromisoformat(f["as_of"]),
                                      datetime.time(0, 0))
    props.created = stamp
    props.modified = stamp
    return prs


def bounds_violations(prs):
    """Every shape inside the slide: left and top non-negative, right and bottom within the
    slide; tables measured by their row heights."""
    problems = []
    width, height = prs.slide_width, prs.slide_height
    for number, sl in enumerate(prs.slides, 1):
        for shape in sl.shapes:
            h = shape.height
            if getattr(shape, "has_table", False) and shape.has_table:
                h = sum(r.height for r in shape.table.rows)
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > width \
                    or shape.top + h > height:
                problems.append("shape %s on page %d sits outside the slide"
                                % (shape.name, number))
    return problems


# ----------------------------------------------------------------- main

def main(argv=None):
    args = parse_args(argv)
    findings = load_findings(args.findings)
    brand_path = core.find_brand(args.brand)
    brand = load_brand(brand_path)
    import_pptx()
    print("%s: brand file: %s" % (PROG, brand_path))
    pal = Palette(brand)
    slides = compose(findings, pal)
    prs = render(slides, findings, brand, pal)
    problems = bounds_violations(prs)
    if problems:
        raise LayoutError("; ".join(problems[:3]))
    buffer = io.BytesIO()
    prs.save(buffer)
    core.write_bytes_atomic(args.out, buffer.getvalue())
    print("%s: wrote %s (%d slides)" % (PROG, args.out, len(slides)))
    return 0


def run_main():
    try:
        return main()
    except InputError as exc:
        sys.stderr.write("%s: %s\n" % (PROG, exc))
        return 1
    except LayoutError as exc:
        sys.stderr.write("%s: layout check failed: %s\n" % (PROG, exc))
        return 4
    except Exception as exc:  # one clean line, never a traceback
        sys.stderr.write("%s: unexpected error: %s\n" % (PROG, exc))
        return 1


if __name__ == "__main__":
    sys.exit(run_main())

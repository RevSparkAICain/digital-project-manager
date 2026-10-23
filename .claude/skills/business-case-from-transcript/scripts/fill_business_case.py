"""Place business case content into a Word template by the anchors a profile names.

Usage:
    python fill_business_case.py --template <blank.docx> --profile <anchors.json>
        --content <content.json> --out <case.docx>

The agent picks the template and profile from the pod config, read from the project's own
file, if present (pods/pm/pm-config.yaml), otherwise from this folder's shared/ copy of the
example. An empty business_case.template means "generate one with make_template.py"; an
empty business_case.profile means references/anchors.example.json.

Profile (JSON):
    {"fields":    {key: "<label in column 1 of a details table row>"},
     "sections":  {key: {"heading": "<heading text>", "placeholder": "<opening words>"}},
     "tables":    {key: {"header": ["<cell text>", ...]}},
     "registers": {"sources": "<title>", "missing": "<title>"}}
Content (JSON, strings only):
    {"fields": {key: str}, "sections": {key: [str, ...]}, "tables": {key: [[str, ...], ...]},
     "sources": [str, ...], "missing": [str, ...]}
fields, sections and tables hold exactly the profile's keys. Every section list has at least
one paragraph, every table row is as wide as its header, sources is never empty and missing
may be empty.

What changes in the document (nothing else is touched):
    Fields     column 2 of the first table row whose column-1 text, without a trailing colon,
               equals the label.
    Sections   under a heading paragraph (style name starting "Heading") whose text equals
               the heading, the first paragraph before the next heading that begins with the
               placeholder words is replaced by the first content paragraph; the rest follow
               it. Each keeps the placeholder's paragraph properties (style, indents,
               spacing); the text goes in as plain runs, so content never looks like a hint.
    Tables     the first table whose header row equals the header. Content rows fill the
               blank data rows in order (blank: no visible text); any overflow goes in as
               copies of the last filled row, inserted straight after it, so a labeled row
               such as a totals line stays last. Each written cell keeps its own properties
               and the style of its first paragraph, and holds only the new text.
    Registers  at the end of the body: a Heading 1 with each register title, then one
               List Bullet paragraph per item, or a "None." paragraph when the list is empty.
               When the template has no Heading 1 or List Bullet style, titles take the style
               of the matched section headings and items take the style of the first replaced
               placeholder with a leading bullet character; one note on stderr names the
               missing styles and the run still exits 0.
    Properties title "Business case: <fields.project>"; author and last modified by
               <fields.prepared_by>; each only when that field is in the profile.
Every comparison collapses whitespace (spaces, tabs, line breaks) and treats typographic
quotes as straight ones. Text inside hyperlinks and simple fields counts as paragraph text.

Order of checks: profile shape, content shape, --out against --template (any spelling of the
same path is refused), python-docx, the template file, then every anchor. All anchors that
cannot be found are named together, as dotted keys (fields.project, sections.summary,
tables.costs). The output is saved only when every anchor was found, through a temporary file
beside it, so a failed run leaves no file and an existing file at --out untouched. The
template itself is only read.

Exit codes: 0 written (the path is printed); 1 a profile, content, template or anchor problem,
or python-docx not installed (one line on stderr); 2 a usage error from argparse.

python-docx is imported inside fill only; the JSON checks need the standard library alone.
"""
import argparse
import copy
import importlib.util
import json
import os
import sys
import tempfile

SKILL_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir))
SHARED_RENDERER = os.path.join(SKILL_DIR, "shared", "md2docx.py")

PROFILE_PARTS = ("fields", "sections", "tables", "registers")
CONTENT_PARTS = ("fields", "sections", "tables", "sources", "missing")
REGISTERS = ("sources", "missing")
TITLE_STYLE, ITEM_STYLE = "Heading 1", "List Bullet"
BULLET = "\u2022 "
EMPTY_REGISTER = "None."

_WML = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_QUOTES = str.maketrans({"\u2018": "'", "\u2019": "'", "\u201a": "'", "\u201b": "'",
                         "\u201c": '"', "\u201d": '"', "\u201e": '"', "\u201f": '"'})


def _w(tag):
    return f"{{{_WML}}}{tag}"


W_P, W_R, W_T, W_PPR = _w("p"), _w("r"), _w("t"), _w("pPr")
W_TBL, W_TR, W_TC, W_TCPR = _w("tbl"), _w("tr"), _w("tc"), _w("tcPr")
W_TRPR, W_SECTPR = _w("trPr"), _w("sectPr")
_SPACERS = (_w("tab"), _w("br"), _w("cr"))


def _load_renderer():
    """Import shared/md2docx.py under a name of its own, for same_path."""
    spec = importlib.util.spec_from_file_location("business_case_fill_md2docx",
                                                  SHARED_RENDERER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


renderer = _load_renderer()


# ---------------------------------------------------------------- JSON checks


def load_json(path):
    """Read a JSON file as UTF-8 (a leading byte order mark is accepted)."""
    with open(path, encoding="utf-8-sig") as fh:
        try:
            return json.load(fh)
        except json.JSONDecodeError as exc:
            raise ValueError(f"{path}: not valid JSON ({exc})") from exc


def _text_value(value):
    return isinstance(value, str) and value.strip() != ""


def check_profile(profile):
    """Return the profile when its shape is right; raise ValueError naming the bad part."""
    if not isinstance(profile, dict):
        raise ValueError("profile must be a JSON object")
    for part in PROFILE_PARTS:
        if not isinstance(profile.get(part), dict):
            raise ValueError(f"profile has no {part} object")
    extra = sorted(set(profile) - set(PROFILE_PARTS))
    if extra:
        raise ValueError(f"unknown profile part: {', '.join(extra)}")
    for key, label in profile["fields"].items():
        if not _text_value(label):
            raise ValueError(f"profile fields.{key} needs the label text")
    for key, entry in profile["sections"].items():
        for name in ("heading", "placeholder"):
            if not isinstance(entry, dict) or not _text_value(entry.get(name)):
                raise ValueError(f"profile sections.{key} needs a {name}")
    for key, entry in profile["tables"].items():
        header = entry.get("header") if isinstance(entry, dict) else None
        if not isinstance(header, list) or not header or not all(map(_text_value, header)):
            raise ValueError(f"profile tables.{key} needs a header list of cell texts")
    for name in REGISTERS:
        if not _text_value(profile["registers"].get(name)):
            raise ValueError(f"profile registers.{name} needs a title")
    return profile


def _string_list(value, where, allow_empty):
    if not isinstance(value, list) or (not value and not allow_empty):
        raise ValueError(f"{where} must be a {'' if allow_empty else 'non-empty '}"
                         f"list of strings")
    for index, item in enumerate(value):
        if not isinstance(item, str):
            raise ValueError(f"{where} item {index + 1} must be a string")


def _same_keys(given, wanted, part):
    if not isinstance(given, dict):
        raise ValueError(f"content {part} must be an object")
    for key in given:
        if key not in wanted:
            raise ValueError(f"unknown content key: {part}.{key}")
    for key in wanted:
        if key not in given:
            raise ValueError(f"content has no {part}.{key}")


def check_content(content, profile):
    """Return the content when it fits the profile; raise ValueError naming the dotted key."""
    if not isinstance(content, dict):
        raise ValueError("content must be a JSON object")
    for key in content:
        if key not in CONTENT_PARTS:
            raise ValueError(f"unknown content key: {key}")
    for key in CONTENT_PARTS:
        if key not in content:
            raise ValueError(f"content has no {key}")
    for part in ("fields", "sections", "tables"):
        _same_keys(content[part], profile[part], part)
    for key, value in content["fields"].items():
        if not isinstance(value, str):
            raise ValueError(f"fields.{key} must be a string")
    for key, paragraphs in content["sections"].items():
        _string_list(paragraphs, f"sections.{key}", allow_empty=False)
    for key, rows in content["tables"].items():
        width = len(profile["tables"][key]["header"])
        if not isinstance(rows, list):
            raise ValueError(f"tables.{key} must be a list of rows")
        for number, row in enumerate(rows, 1):
            _string_list(row, f"tables.{key} row {number}", allow_empty=True)
            if len(row) != width:
                raise ValueError(f"tables.{key} row {number} has {len(row)} cells; "
                                 f"the header has {width}")
    _string_list(content["sources"], "sources", allow_empty=False)
    _string_list(content["missing"], "missing", allow_empty=True)
    return content


# ---------------------------------------------------------------- text comparison


def _norm(text):
    """Straight quotes, whitespace runs collapsed to one space, ends trimmed."""
    return " ".join(str(text).translate(_QUOTES).split())


def _label(text):
    return _norm(text).rstrip(":").rstrip()


def _text(element):
    """Visible text of a paragraph, cell or row: every w:t in order, hyperlinks and fields
    included; tabs, breaks and paragraph ends count as spaces."""
    parts = []
    for node in element.iter():
        if node.tag == W_T:
            parts.append(node.text or "")
        elif node.tag in _SPACERS and node.getparent().tag == W_R:
            parts.append(" ")
        elif node.tag == W_P:
            parts.append(" ")
    return "".join(parts)


# ---------------------------------------------------------------- locating anchors


def _find_field(body, label):
    wanted = _label(label)
    for row in body.iter(W_TR):
        cells = row.findall(W_TC)
        if len(cells) >= 2 and _label(_text(cells[0])) == wanted:
            return cells[1]
    return None


def _find_section(paragraphs, heading, placeholder):
    """Return (heading paragraph, placeholder paragraph) or None. paragraphs is the body's
    paragraph list; a section runs from its heading to the next heading."""
    heading, placeholder = _norm(heading), _norm(placeholder)
    current = None
    for paragraph in paragraphs:
        if (getattr(paragraph.style, "name", "") or "").startswith("Heading"):
            current = paragraph if _norm(_text(paragraph._p)) == heading else None
            continue
        if current is not None and _norm(_text(paragraph._p)).startswith(placeholder):
            return current, paragraph
    return None


def _find_table(body, header, claimed):
    """First table with this header row that no earlier profile table took. claimed holds the
    elements themselves: lxml can hand a freed element's id() to a new one."""
    wanted = [_norm(cell) for cell in header]
    for table in body.iter(W_TBL):
        if any(table is taken for taken in claimed):
            continue
        rows = table.findall(W_TR)
        if rows and [_norm(_text(cell)) for cell in rows[0].findall(W_TC)] == wanted:
            return table
    return None


def _table_layout(table, count, width):
    """Return (rows to fill, row to copy for overflow, overflow count), or None when a row
    that would be written has fewer cells than the header."""
    data = table.findall(W_TR)[1:]
    slots = [row for row in data if _norm(_text(row)) == ""]
    used = slots[:count]
    overflow = count - len(used)
    base = used[-1] if used else (data[-1] if data else table.findall(W_TR)[0])
    for row in used + ([base] if overflow else []):
        if len(row.findall(W_TC)) < width:
            return None
    return used, base, overflow


# ---------------------------------------------------------------- writing


def _plain_run(paragraph_element, text):
    """Append one run without character properties; newlines and tabs become breaks and tabs."""
    from docx.text.paragraph import Paragraph

    Paragraph(paragraph_element, None).add_run(text)


def _strip_to_properties(paragraph_element):
    for child in list(paragraph_element):
        if child.tag != W_PPR:
            paragraph_element.remove(child)


def _bare_copy(paragraph_element):
    """A new w:p holding a copy of the paragraph properties, without any section break."""
    new = copy.deepcopy(paragraph_element)
    _strip_to_properties(new)
    properties = new.find(W_PPR)
    if properties is not None:
        for brk in properties.findall(W_SECTPR):
            properties.remove(brk)
    return new


def _replace_paragraph(paragraph_element, texts):
    _strip_to_properties(paragraph_element)
    _plain_run(paragraph_element, texts[0])
    last = paragraph_element
    for text in texts[1:]:
        new = _bare_copy(paragraph_element)
        last.addnext(new)
        _plain_run(new, text)
        last = new
    properties = paragraph_element.find(W_PPR)
    if last is not paragraph_element and properties is not None:
        for brk in properties.findall(W_SECTPR):
            last.find(W_PPR).append(brk)


def _set_cell(cell, text):
    """Leave one paragraph in the cell, keep its properties, and give it the text."""
    paragraphs = cell.findall(W_P)
    if not paragraphs:
        from docx.oxml import OxmlElement

        paragraphs = [OxmlElement("w:p")]
        cell.append(paragraphs[0])
    for child in list(cell):
        if child.tag != W_TCPR and child is not paragraphs[0]:
            cell.remove(child)
    _strip_to_properties(paragraphs[0])
    _plain_run(paragraphs[0], text)


def _set_row(row, values):
    for cell, value in zip(row.findall(W_TC), values):
        _set_cell(cell, value)


def _copy_row(base):
    new = copy.deepcopy(base)
    properties = new.find(W_TRPR)
    if properties is not None:
        for flag in properties.findall(_w("tblHeader")):
            properties.remove(flag)
    return new


def _fill_table(layout, rows):
    used, base, _ = layout
    for row, values in zip(used, rows):
        _set_row(row, values)
    anchor = base if used else base.getparent().findall(W_TR)[-1]
    for values in rows[len(used):]:
        new = _copy_row(base)
        anchor.addnext(new)
        _set_row(new, values)
        anchor = new


def _style_or_none(document, name):
    try:
        return document.styles[name]
    except KeyError:
        return None


def _append_registers(document, profile, content, heading_style, item_style):
    """Add the two registers; return the names of styles the template lacks."""
    title = _style_or_none(document, TITLE_STYLE)
    bullet = _style_or_none(document, ITEM_STYLE)
    lacking = [name for name, found in ((TITLE_STYLE, title), (ITEM_STYLE, bullet))
               if found is None]
    for name in REGISTERS:
        document.add_paragraph(profile["registers"][name], style=title or heading_style)
        items = content[name]
        if not items:
            document.add_paragraph(EMPTY_REGISTER)
        for item in items:
            if bullet is not None:
                document.add_paragraph(item, style=bullet)
            else:
                document.add_paragraph(BULLET + item, style=item_style)
    return lacking


def _save_beside(document, out):
    """Save through a temporary file in the target folder, then move it into place."""
    folder = os.path.dirname(os.path.abspath(out))
    handle, temporary = tempfile.mkstemp(suffix=".docx", dir=folder)
    os.close(handle)
    try:
        document.save(temporary)
        os.replace(temporary, out)
    except BaseException:
        if os.path.exists(temporary):
            os.remove(temporary)
        raise


def _refuse_template_as_out(template, out):
    if renderer.same_path(template, out):
        raise ValueError(f"--out names the template ({template}); choose another output path")


def fill(template, profile, content, out):
    """Fill a copy of template with content by the profile's anchors, save it to out and
    return out. Raises ValueError naming every anchor that was not found."""
    profile = check_profile(profile)
    content = check_content(content, profile)
    _refuse_template_as_out(template, out)
    from docx import Document

    if not os.path.isfile(template):
        raise ValueError(f"template not found: {template}")
    try:
        document = Document(template)
    except Exception as exc:  # python-docx raises several types for a file that is not .docx
        raise ValueError(f"template is not a readable Word file: {template} "
                         f"({type(exc).__name__})") from exc
    body = document.element.body
    paragraphs = document.paragraphs
    missing, problems = [], []

    fields = {}
    for key, label in profile["fields"].items():
        cell = _find_field(body, label)
        if cell is None:
            missing.append(f"fields.{key}")
        else:
            fields[key] = cell

    sections = {}
    for key, entry in profile["sections"].items():
        found = _find_section(paragraphs, entry["heading"], entry["placeholder"])
        if found is None:
            missing.append(f"sections.{key}")
        else:
            sections[key] = found

    tables, claimed = {}, []
    for key, entry in profile["tables"].items():
        table = _find_table(body, entry["header"], claimed)
        if table is None:
            missing.append(f"tables.{key}")
            continue
        claimed.append(table)
        layout = _table_layout(table, len(content["tables"][key]), len(entry["header"]))
        if layout is None:
            problems.append(f"tables.{key} has a row with fewer cells than its header "
                            f"(merged cells?)")
        else:
            tables[key] = layout

    if missing:
        problems.insert(0, f"anchors not found in {template}: {', '.join(missing)}")
    if problems:
        raise ValueError("; ".join(problems))

    for key, cell in fields.items():
        _set_cell(cell, content["fields"][key])
    heading_style = item_style = None
    for key, (heading, placeholder) in sections.items():
        if heading_style is None:
            heading_style, item_style = heading.style, placeholder.style
        _replace_paragraph(placeholder._p, content["sections"][key])
    for key, layout in tables.items():
        _fill_table(layout, content["tables"][key])
    lacking = _append_registers(document, profile, content, heading_style, item_style)

    props = document.core_properties
    if "project" in content["fields"]:
        props.title = f"Business case: {content['fields']['project']}"
    if "prepared_by" in content["fields"]:
        props.author = content["fields"]["prepared_by"]
        props.last_modified_by = content["fields"]["prepared_by"]

    _save_beside(document, out)
    if lacking:
        print(f"note: {os.path.basename(template)} has no {' or '.join(lacking)} style; "
              f"the registers use the section heading and placeholder styles instead",
              file=sys.stderr)
    return out


def _one_line(message):
    return " ".join(str(message).split())


def main(argv=None):
    parser = argparse.ArgumentParser(prog="fill_business_case.py",
                                     description="Fill a business case template from a "
                                                 "content file by a profile's anchors.")
    parser.add_argument("--template", required=True, help="the blank .docx to fill (only read)")
    parser.add_argument("--profile", required=True, help="the anchors profile JSON")
    parser.add_argument("--content", required=True, help="the content JSON")
    parser.add_argument("--out", required=True, help="the .docx to write")
    args = parser.parse_args(argv)
    try:
        profile = check_profile(load_json(args.profile))
        content = check_content(load_json(args.content), profile)
        _refuse_template_as_out(args.template, args.out)
        fill(args.template, profile, content, args.out)
    except ImportError:
        print("not filled: python-docx is not installed; install it with: "
              "pip install python-docx", file=sys.stderr)
        return 1
    except (OSError, ValueError) as exc:
        print(_one_line(f"not filled: {exc}"), file=sys.stderr)
        return 1
    print(args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())

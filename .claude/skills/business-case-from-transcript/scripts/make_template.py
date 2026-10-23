"""Build a blank, branded business case template (.docx) for fill_business_case.py to fill.

Usage:
    python make_template.py --out <template.docx> [--brand <brand.json>]

Run it when the team has no template of its own. The agent reads business_case.template from
the pod config, the project's own file, if present (pods/pm/pm-config.yaml), otherwise this
folder's shared/ copy of the example; an empty value means run this script. The script itself
never reads the config.

Layout, top to bottom:
    1. A Title paragraph, "Business case".
    2. A two-column Table Grid table: the DETAILS labels (Project, Sponsor, Prepared by, Date)
       in column 1, column 2 left empty.
    3. For each entry in SECTIONS, in order: a Heading 1 paragraph with the heading, then the
       placeholder sentence as a Normal paragraph. The four sections named in TABLES are each
       followed by a Table Grid table with a bold header row and one empty data row.
references/anchors.example.json describes this layout for the fill script: each label, each
heading with the opening words of its placeholder, and each header row.

Styling comes from shared/md2docx.py, this folder's copy of the notes renderer, so the case
and the meeting notes share one look. Brand file lookup, first hit wins:
    1. --brand
    2. the BRAND_JSON environment variable
    3. the nearest brand/brand.json in the renderer's folder or a folder above it, up to and
       including the project root
    4. shared/brand.json in this skill folder
A path named by 1 or 2 that does not exist is an error. Styles are branded before any content
is added; the theme colors and fonts are rewritten from the brand file just before saving.

Repeatable output: author and last modified by are GENERATOR, title is
"Business case template", created and modified are 2030-01-01 00:00 and revision is 1, so two
runs with the same brand file give the same document, styles and core property XML.

Exit codes: 0 written (the path is printed); 1 a brand or file problem, or python-docx is not
installed (one line on stderr); 2 a usage error from argparse, such as a missing --out.

python-docx is imported inside build_template only, so the argument checks run without it.
"""
import argparse
import datetime
import importlib.util
import json
import os
import sys

SKILL_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir))
SHARED_RENDERER = os.path.join(SKILL_DIR, "shared", "md2docx.py")

GENERATOR = "ACME Co. demo generator"
TEMPLATE_TITLE = "Business case template"
FIXED_TIME = datetime.datetime(2030, 1, 1, 0, 0)

DETAILS = ("Project", "Sponsor", "Prepared by", "Date")
SECTIONS = [
    ("Summary", "Summarize the request, the recommended option and the decision being asked "
                "for."),
    ("Problem or opportunity", "Describe the problem or opportunity, who it affects and what "
                               "it costs today."),
    ("Options considered", "Describe each option considered, including making no change."),
    ("Recommended option", "Name the recommended option and the reasons it was chosen."),
    ("Costs and benefits", "List each cost and benefit with its amount and timing."),
    ("Risks", "List the main risks of the recommended option and how each is handled."),
    ("Timeline", "List the milestones with target dates and owners."),
    ("Approval", "Record who approves the case and their decision."),
]
TABLES = {
    "Costs and benefits": ("Item", "Type", "Amount", "Timing"),
    "Risks": ("Risk", "Likelihood", "Impact", "Score", "Response", "Owner"),
    "Timeline": ("Milestone", "Target date", "Owner"),
    "Approval": ("Role", "Name", "Decision", "Date"),
}


def _load_renderer():
    """Import shared/md2docx.py under a name of its own. Importing it needs only the standard
    library, so the usage checks still run where python-docx is missing."""
    spec = importlib.util.spec_from_file_location("business_case_template_md2docx",
                                                  SHARED_RENDERER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


renderer = _load_renderer()


def load_brand(brand=None):
    """Find the brand file by the lookup order above; return its (colors, fonts)."""
    path = renderer.locate_brand(brand)
    with open(path, encoding="utf-8") as fh:
        try:
            data = json.load(fh)
        except json.JSONDecodeError as exc:
            raise ValueError(f"{path}: not valid JSON ({exc})") from exc
    return data["colors"], data["fonts"]


def _grid(document, rows, cols):
    table = document.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    return table


def build_template(out, brand=None):
    """Write the blank template to out and return out. brand is an optional brand.json path."""
    colors, fonts = load_brand(brand)
    from docx import Document

    document = Document()
    body = document.element.body
    for stray in body.findall(renderer._w("p")):
        body.remove(stray)
    renderer._brand_styles(document, colors, fonts)

    document.add_paragraph("Business case", style="Title")
    details = _grid(document, len(DETAILS), 2)
    for row, label in zip(details.rows, DETAILS):
        row.cells[0].text = label

    for heading, placeholder in SECTIONS:
        document.add_paragraph(heading, style="Heading 1")
        document.add_paragraph(placeholder, style="Normal")
        header = TABLES.get(heading)
        if header:
            table = _grid(document, 2, len(header))
            for cell, text in zip(table.rows[0].cells, header):
                cell.paragraphs[0].add_run(text).bold = True

    props = document.core_properties
    props.author = GENERATOR
    props.last_modified_by = GENERATOR
    props.title = TEMPLATE_TITLE
    props.comments = ""
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.revision = 1

    renderer._apply_theme(document, colors, fonts)
    renderer._sync_legacy_styles(document)
    document.save(out)
    return out


def _one_line(message):
    return " ".join(str(message).split())


def main(argv=None):
    parser = argparse.ArgumentParser(prog="make_template.py",
                                     description="Build a blank, branded business case "
                                                 "template.")
    parser.add_argument("--out", required=True, help="the .docx template to write")
    parser.add_argument("--brand", help="brand.json to use (see the lookup order)")
    args = parser.parse_args(argv)
    try:
        build_template(args.out, brand=args.brand)
    except ImportError:
        print("template not written: python-docx is not installed; install it with: "
              "pip install python-docx", file=sys.stderr)
        return 1
    except KeyError as exc:
        print(_one_line(f"template not written: brand or theme entry missing: "
                        f"{exc.args[0] if exc.args else exc}"), file=sys.stderr)
        return 1
    except (OSError, ValueError) as exc:
        print(_one_line(f"template not written: {exc}"), file=sys.stderr)
        return 1
    print(args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())

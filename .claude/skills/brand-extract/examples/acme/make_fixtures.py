"""Build the binary ACME demo collateral from code: a two-slide deck, a one-page brand sheet
PDF and a raster logo. ACME Co. is fictional. The site, its CSS and logo.svg are committed as
text; these binaries are generated so no binary fixture is stored in git.

    python make_fixtures.py --out <folder>
"""
import argparse
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.normpath(os.path.join(HERE, "..", "..", "scripts"))
if SCRIPTS not in sys.path:
    sys.path.insert(0, SCRIPTS)

import office_theme  # noqa: E402

RED, SAND, GRAY = "#C8102E", "#F5EBDD", "#3A3F44"
ORANGE, TEAL = "#E87722", "#00797E"
SLATE, MIST, STEEL = "#5B6066", "#BFC2C4", "#8A8D8F"
HEADING_FONT, BODY_FONT = "Inter", "Source Serif 4"
DECK, SHEET, LOGO = "acme-deck.pptx", "acme-brand-sheet.pdf", "acme-logo.png"


def _rgb01(hexv):
    return tuple(int(hexv[i:i + 2], 16) / 255 for i in (1, 3, 5))


def make_deck(path):
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "ACME Co."
    slide.placeholders[1].text = "Quarterly update (fictional demo deck)"
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "What we ship"
    slide.placeholders[1].text = "Tools, parts and same-day delivery."
    props = prs.core_properties
    props.author, props.last_modified_by = "ACME Co.", "make_fixtures"
    props.comments, props.title = "Fictional demo collateral", "ACME Co. demo deck"
    prs.save(path)
    office_theme.apply_theme(path, {
        "dk1": GRAY, "lt1": SAND, "dk2": GRAY, "lt2": SAND,
        "accent1": RED, "accent2": ORANGE, "accent3": TEAL,
        "accent4": STEEL, "accent5": MIST, "accent6": SLATE,
        "hlink": RED, "folHlink": SLATE}, {"major": HEADING_FONT, "minor": BODY_FONT})


def make_sheet(path):
    import fitz
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.draw_rect(fitz.Rect(0, 0, 612, 792), color=None, fill=_rgb01(SAND), width=0)
    page.draw_rect(fitz.Rect(0, 0, 612, 110), color=None, fill=_rgb01(RED), width=0)
    page.draw_rect(fitz.Rect(0, 110, 612, 118), color=None, fill=_rgb01(ORANGE), width=0)
    page.insert_text((48, 70), "ACME Co. brand sheet", fontsize=30, fontname="helv",
                     color=_rgb01(SAND))
    swatches = [(RED, "Roadrunner Red"), (ORANGE, "Canyon Orange"), (TEAL, "Mesa Teal"),
                (GRAY, "Anvil Gray")]
    for i, (hexv, label) in enumerate(swatches):
        x = 48 + i * 132
        page.draw_rect(fitz.Rect(x, 170, x + 110, 250), color=None, fill=_rgb01(hexv), width=0)
        page.insert_text((x, 270), label, fontsize=11, fontname="helv", color=_rgb01(GRAY))
        page.insert_text((x, 286), hexv, fontsize=11, fontname="helv", color=_rgb01(GRAY))
    notes = ["Headings: Inter. Body text: Source Serif 4.",
             "Background: Desert Sand " + SAND + ".",
             "ACME Co. is a fictional company used for demonstration."]
    for j, text in enumerate(notes):
        page.insert_text((48, 340 + j * 20), text, fontsize=12, fontname="helv",
                         color=_rgb01(GRAY))
    doc.set_metadata({"title": "ACME Co. brand sheet", "author": "ACME Co.",
                      "creator": "make_fixtures", "producer": "make_fixtures"})
    doc.save(path)
    doc.close()


def make_logo_png(path):
    from PIL import Image, ImageDraw
    img = Image.new("RGBA", (240, 80), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, 79, 79), fill=RED)
    draw.polygon([(16, 60), (40, 18), (64, 60)], fill=SAND)
    draw.rectangle((96, 30, 230, 50), fill=GRAY)
    img.save(path)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Build the ACME demo deck, brand sheet and logo.")
    ap.add_argument("--out", required=True)
    args = ap.parse_args(argv)
    os.makedirs(args.out, exist_ok=True)
    for name, maker in ((DECK, make_deck), (SHEET, make_sheet), (LOGO, make_logo_png)):
        maker(os.path.join(args.out, name))
        print("wrote %s" % name)
    return 0


if __name__ == "__main__":
    sys.exit(main())
